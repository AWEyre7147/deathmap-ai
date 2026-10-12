"""Read-only verification of native receipts, projection and the final workbook.

Every exported cell is checked against canonical metadata; no source or workbook
is repaired here. Publication associations are checked structurally, not judged
as qualifying screens. A validation receipt is written only after all checks pass.
"""
import hashlib
import json
from pathlib import Path
import openpyxl

base=Path(r'C:\Users\aweyr\Documents\Repositories\deathmap-ai')
folder=base/'outputs/arrayexpress-orcs-20261011-S-network'
load=lambda p:json.loads(p.read_text(encoding='utf-8'))
database=load(folder/'database.json')
index=load(base/'data/orcs/publication-index.json')
assert {p['publication']['PUBLICATION_ID'] for p in database['publications']}=={p['PUBLICATION_ID'] for p in index}
assert len(database['publications'])==len(index)==418
assert hashlib.sha256((base/'data/orcs/publication-index.json').read_bytes()).hexdigest()==database['input_sha256']
manifest=load(base/'data/templates/BioStudies-ArrayExpress-Enrichment.manifest.json')
assert hashlib.sha256((base/manifest['template']).read_bytes()).hexdigest()==manifest['sha256']
receipts=load(folder/'requests.json');verified=0
for receipt in receipts:
    if receipt.get('file'):
        body=(folder/receipt['file']).read_bytes()
        assert len(body)==receipt['bytes']
        assert hashlib.sha256(body).hexdigest()==receipt['sha256']
        verified+=1
for accession,study in database['studies'].items():
    assert study['native']['accno']==accession
    assert (folder/study['source']).exists()
    for metadata in study['metadata_files']:
        assert metadata['kind'] in ('idf','sdrf') and (folder/metadata['source']).exists()
confirmed={a for a,s in database['studies'].items() if s['arrayexpress_membership_confirmed']}
assert confirmed=={r['Study accession'] for r in database['rows']['Studies']}
workbook=openpyxl.load_workbook(folder/'BioStudies-ArrayExpress-Enrichment.xlsx')
assert workbook.sheetnames==[s['name'] for s in manifest['sheets']]
compared=0
for info in manifest['sheets']:
    sheet=workbook[info['name']]
    expected=[[row.get(c['source_key']) for c in info['columns']] for row in database['rows'][info['source_table']]]
    assert sheet.max_row==len(expected)+1 and sheet.max_column==len(info['headers'])
    assert list(next(sheet.values))==info['headers']
    for number,(actual,wanted) in enumerate(zip(sheet.iter_rows(min_row=2,values_only=True),expected),2):
        for column,(value,target) in enumerate(zip(actual,wanted),1):
            if target=='':target=None
            assert value==target,(info['name'],number,column,repr(value),repr(target))
            compared+=1
workbook.close()
report={'status':'passed','publications':418,'native_responses_hash_verified':verified,
        'exported_cells_compared':compared,'template_sha256':manifest['sha256'],
        'workbook_sha256':hashlib.sha256((folder/'BioStudies-ArrayExpress-Enrichment.xlsx').read_bytes()).hexdigest(),
        'coverage_claim':'retrieval accounting only; scientific screen attribution unassessed'}
(folder/'verification.json').write_text(json.dumps(report,indent=2),encoding='utf-8')
print(json.dumps(report))
