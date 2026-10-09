"""Refresh the three existing conditions after shared identifier enrichment.

All new runs must export successfully before older runs are archived and removed.
Archive copies are verified byte-for-byte, including any owner reviewer edits.
This explicit maintenance script keeps one active run per authorized condition.
"""
import argparse,hashlib,json,shutil
from pathlib import Path
from datetime import datetime,timezone
from deathmap_ai.orcs_crispr_pilot import run_and_export


def refresh(root,archive):
    """Generate three new runs, verify values, then preserve and retire old runs."""
    from openpyxl import load_workbook
    root=Path(root).resolve();base=root/'outputs/orcs'
    archive=Path(archive).resolve()/'publication-identifiers'/datetime.now(timezone.utc).strftime('%Y%m%dT%H%M%S%fZ')/'outputs'
    names=['cancer-cell-crispr-knockout-v01','cancer-cell-crispr-knockout-v02','orcs-crispr-biological-classes-pilot-v01']
    old={name:[p for p in (base/name).iterdir() if p.is_dir()] for name in names};new={}
    for name in names:
        out=run_and_export(root,profile_path=f'searches/orcs/{name}/profile.json');new[name]=out
        projection=json.loads((out/'projection.json').read_text(encoding='utf-8'))
        book=load_workbook(out/'workbook-authored.xlsx')
        for sheet,rows in projection['sheets'].items():
            headers=[c.value for c in book[sheet][1]]
            for number,row in enumerate(rows,2):
                for column,key in enumerate(headers,1):
                    actual=book[sheet].cell(number,column).value
                    if actual!=row.get(key) and not (actual is None and row.get(key)==''):raise ValueError('Workbook projection mismatch')
        book.close();print(out,flush=True)
    archived=[]
    for name,paths in old.items():
        for source in paths:
            destination=archive/name/source.name
            if not source.resolve().is_relative_to(base.resolve()) or not destination.resolve().is_relative_to(archive):raise ValueError('Unsafe archive target')
            shutil.copytree(source,destination)
            for path in source.rglob('*'):
                if path.is_file():
                    digest=hashlib.sha256(path.read_bytes()).hexdigest()
                    if hashlib.sha256((destination/path.relative_to(source)).read_bytes()).hexdigest()!=digest:raise ValueError('Archive mismatch')
                    archived.append({'original':path.relative_to(root).as_posix(),'sha256':digest})
            shutil.rmtree(source)
    report={'runs':{name:path.relative_to(root).as_posix() for name,path in new.items()},'archive':str(archive),'archived_files':archived,'coverage':{}}
    for name,path in new.items():
        rows=json.loads((path/'projection.json').read_text(encoding='utf-8'))['sheets']['Publications']
        report['coverage'][name]={'publications':len(rows),**{k:sum(bool(row.get(k)) for row in rows) for k in ['pmid','doi','pmcid']}}
    (root/'logs/orcs-identifier-refresh.json').write_text(json.dumps(report,indent=2)+'\n',encoding='utf-8')
    lines=['# Curated ORCS outputs','','One identifier-enriched run per condition. Older workbooks are preserved in the verified external archive; no reviewed workbook was rewritten.','','| Search | Publications | PMID | DOI | PMCID |','|---|---:|---:|---:|---:|']
    for name,counts in report['coverage'].items():
        relative=new[name].relative_to(base).as_posix()
        lines.append(f'| [{name}]({relative}/workbook-authored.xlsx) | '+ ' | '.join(str(counts[k]) for k in ['publications','pmid','doi','pmcid'])+' |')
    lines+=['','Exact NCBI publication identifier lookups enrich the shared catalog before searches. Missing identifiers remain blank. Sources-sheet example counts are illustrative; use summary.json for actual screen counts. Publication associations are not screen-dataset attribution.']
    (base/'README.md').write_text('\n'.join(lines)+'\n',encoding='utf-8')
    print(json.dumps(report['coverage'],indent=2))


if __name__=='__main__':
    p=argparse.ArgumentParser(description=__doc__);p.add_argument('--root',type=Path,default=Path(__file__).resolve().parents[1]);p.add_argument('--archive',type=Path,required=True);args=p.parse_args();refresh(args.root,args.archive)
