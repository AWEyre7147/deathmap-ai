"""Validate a metadata-only ORCS package and prepare explicit enrichment seeds.

Read workbook values and OOXML independently of its writer. Require reviewer
blanks, run-local evidence linkage and no template-derived unsupported entities.
The resulting seed manifest is input preparation, not repository retrieval or
proof that a publication-associated deposit contains a qualifying screen.
"""
import hashlib
import json
from pathlib import Path
import re
import sys
import xml.etree.ElementTree as E
import zipfile
from openpyxl import load_workbook

ROOT=Path(__file__).resolve().parents[1]
N='http://schemas.openxmlformats.org/spreadsheetml/2006/main'
Q=lambda tag:f'{{{N}}}{tag}'
ERRORS={'#NULL!','#DIV/0!','#VALUE!','#REF!','#NAME?','#NUM!','#N/A','#GETTING_DATA'}


def verify_package(path):
    """Verify all XLSX XML parts, cached cell types and index bounds."""
    with zipfile.ZipFile(path) as z:
        if z.testzip():raise ValueError('Invalid ZIP member checksum')
        parsed={n:E.fromstring(z.read(n)) for n in z.namelist() if n.endswith('.xml') or n.endswith('.rels')}
        styles=len(parsed['xl/styles.xml'].find(Q('cellXfs')))
        strings=len(parsed['xl/sharedStrings.xml']) if 'xl/sharedStrings.xml' in parsed else 0
        checked=0
        for name,root in parsed.items():
            if not name.startswith('xl/worksheets/') or not name.endswith('.xml'):continue
            seen=set()
            for c in root.iter(Q('c')):
                addr=c.get('r')
                if addr in seen:raise ValueError(f'Duplicate cell: {name}/{addr}')
                seen.add(addr);checked+=1
                if int(c.get('s','0'))>=styles:raise ValueError('Invalid style reference')
                value=c.find(Q('v'))
                if c.get('t')=='s' and (value is None or not 0<=int(value.text)<strings):raise ValueError('Invalid shared-string reference')
                if c.get('t')=='e':
                    if value is None or value.text not in ERRORS:raise ValueError('Invalid Excel error cache')
                    raise ValueError(f'Formula error: {name}/{addr}/{value.text}')
                if c.find(Q('f')) is not None and (c.find(Q('f')).text or '').startswith('HYPERLINK('):
                    if c.get('t')!='str' or value is None or not value.text:raise ValueError('Missing literal hyperlink cache')
        return {'xml_parts':len(parsed),'cells_checked':checked,'invalid_error_values':0}


def verify(run):
    """Check a completed run, write its review receipt and enrichment seed list."""
    run=Path(run).resolve()
    p=json.loads((run/'projection.json').read_text(encoding='utf-8'))
    ledger=[json.loads(line) for line in (run/'evidence-ledger.jsonl').read_text(encoding='utf-8').splitlines()]
    expected=[f'EVI-{i:05d}' for i in range(1,len(ledger)+1)]
    if [e['evidence_id'] for e in ledger]!=expected:raise ValueError('Noncontiguous ledger IDs')
    if [r['evidence_id'] for r in p['sheets']['Evidence']]!=expected:raise ValueError('Projection/ledger IDs differ')
    if (ROOT/'logs/release-preparation/reviewed-workbook-preservation.json').exists():
        protected=json.loads((ROOT/'logs/release-preparation/reviewed-workbook-preservation.json').read_text(encoding='utf-8'))
        if any((ROOT/r['path']).parent==run for r in protected['reviewed_workbooks']):
            raise ValueError('Owner-reviewed package: do not apply fresh-export blank-field checks or overwrite its receipts')
    receipts={'workbook-authored.xlsx':verify_package(run/'workbook-authored.xlsx')}
    for name in receipts:
        w=load_workbook(run/name,data_only=False)
        for sheet,rows in p['sheets'].items():
            values=list(w[sheet].values);headers=list(values[0])
            actual=[dict(zip(headers,row)) for row in values[1:] if any(v is not None for v in row)]
            if len(actual)!=len(rows):raise ValueError(f'Row count mismatch: {name}/{sheet}: {len(actual)} != {len(rows)}')
            for field in ['reviewer_decision','reviewer_notes']:
                if sheet in {'Publications','Screens'} and field in headers:raise ValueError('Retired review column present')
            for real,source in zip(actual,rows):
                for field,value in source.items():
                    if field=='source_link':
                        if real[field]!=value:raise ValueError('Native source link differs')
                        cell=w[sheet].cell(actual.index(real)+2,headers.index(field)+1)
                        if cell.hyperlink is None or cell.hyperlink.target!=value:raise ValueError('Native hyperlink target differs')
                    elif (real.get(field) or None)!=(value or None):raise ValueError(f'Workbook value differs: {sheet}/{field}')
                if sheet in {'Publications','Screens'}:
                    notes='publication_notes' if sheet=='Publications' else 'screen_notes'
                    if real.get(notes) not in (None,''):raise ValueError('Reviewer notes are not blank')
                    link='evidence_id_link' if sheet=='Publications' else 'source_evidence_ids'
                    if any(eid not in expected for eid in (real.get(link) or '').split('; ') if eid):raise ValueError('Dangling evidence link')
        if any(str(v).strip().lower()=='cell line' for row in w['Definitions'].values for v in row if v is not None):raise ValueError('Cell line definition retained')
        w.close()
    manifest=json.loads((run/'run_manifest.json').read_text(encoding='utf-8'))
    for relative,digest in manifest['input_hashes'].items():
        if hashlib.sha256((ROOT/relative).read_bytes()).hexdigest()!=digest:raise ValueError(f'Input changed: {relative}')
    pubs={m['publication_id']:m['PUBLICATION_ID'] for m in p['native_mapping']['publications']}
    seeds=[]
    for pub in p['sheets']['Publications']:
        screens=[s for s in p['sheets']['Screens'] if s.get('publication_id')==pub['publication_id']]
        seeds.append({'publication_id':pub['publication_id'],'orcs_publication_id':pubs[pub['publication_id']],
            'pmid':pub.get('pmid'),'doi':pub.get('doi'),'title_reported':pub.get('title_original'),
            'evidence_ids':pub.get('evidence_id_link'),'screens':[{'screen_id':s['screen_id'],**p['roles'][s['screen_id']]} for s in screens],
            'repository_accessions':[],'dataset_attribution_status':'not_assessed'})
    (run/'enrichment-seeds.json').write_text(json.dumps({'run_id':manifest['run_id'],'profile_id':manifest['profile_id'],
        'seed_route':'ORCS publication/screen metadata','publications':seeds,'repository_requests':0},indent=2)+'\n',encoding='utf-8')
    receipt={'workbooks':receipts,'reviewer_fields_verified':True,'evidence_links_verified':True,
        'unsupported_entity_examples_removed':True,'source_input_hashes_unchanged':True,
        'native_excel_open_test':'not performed','network_requests':0,'evidence_count':len(ledger),'publication_seeds':len(seeds)}
    (run/'verification.json').write_text(json.dumps(receipt,indent=2)+'\n',encoding='utf-8')
    manifest['status']='metadata_prepared_for_repository_enrichment'
    manifest['workbook_validation']='verification.json';manifest['workbook_export_provenance']='workbook-export-manifest.json'
    (run/'run_manifest.json').write_text(json.dumps(manifest,indent=2)+'\n',encoding='utf-8')
    print(f'{run.name}: {len(seeds)} publication seeds, {len(ledger)} evidence rows; authored workbook structurally verified (not an Excel opening test)')


if __name__=='__main__':verify(sys.argv[1])
