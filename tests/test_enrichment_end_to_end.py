"""Small offline fixtures exercise registered export and complete synthesis.

No benchmark identifiers, live network, large source collection or Codex runtime
is required. Empty repository results and missing PMIDs remain legitimate data.
"""
import argparse
import hashlib
import json
import io
import gzip
from contextlib import redirect_stdout
from pathlib import Path
import shutil
import tempfile
import threading
import unittest
from unittest.mock import patch

import openpyxl
from deathmap_ai.enrichment_commands import publication_registry, run_stage
from deathmap_ai.enrichment_export import export_registered
from deathmap_ai.pubmed_enrichment import Retrieval, API

ROOT = Path(__file__).resolve().parents[1]


class EndToEndTests(unittest.TestCase):
    def test_registry_preserves_keys_and_missing_pmid(self):
        pubs = [{'PUBLICATION_ID': '1', 'PMID': '123'}, {'PUBLICATION_ID': '2', 'PMID': None}]
        old = [{'Publication index': 'PUB-0042', 'ORCS Publication ID': '1', 'PMID': '123'}]
        result = publication_registry(pubs, old)
        self.assertEqual([p['Publication index'] for p in result], ['PUB-0042', 'PUB-0043'])
        self.assertIsNone(result[1]['PMID'])
        self.assertEqual(publication_registry(pubs, result), result)
        with self.assertRaises(ValueError):
            publication_registry(pubs + [pubs[0]], old)

    def test_literal_export_and_empty_sheets(self):
        with tempfile.TemporaryDirectory(dir=ROOT / 'outputs') as t:
            template = ROOT / 'data/templates/Enrichment-Synthesis.xlsx'
            registration = json.loads(template.with_suffix('.manifest.json').read_text(encoding='utf-8'))
            tables = {s['name']: [] for s in registration['sheets']}
            tables['Publications'] = [{'Publication Index': 'PUB-0001', 'PMID': '123',
                'Publication Title': '=literal metadata', 'Publication Date': '2020-01-02'}]
            out = Path(t) / 'export.xlsx'
            receipt = export_registered(template, out, tables)
            self.assertEqual(receipt['sheet_rows']['Publications'], 1)
            book = openpyxl.load_workbook(out)
            self.assertEqual(book['Publications']['E2'].value, '=literal metadata')
            self.assertEqual(book['Publications']['E2'].data_type, 's')
            self.assertEqual(book['Publications']['H2'].value, '2020-01-02')
            book.close()
            with self.assertRaises(FileExistsError):
                export_registered(template, out, tables)

    def test_oversized_membership_is_lossless_and_explicit(self):
        with tempfile.TemporaryDirectory(dir=ROOT/'outputs') as t:
            template = ROOT/'data/templates/Enrichment-Synthesis.xlsx'
            manifest = json.loads(template.with_suffix('.manifest.json').read_text(encoding='utf-8'))
            tables = {s['name']: [] for s in manifest['sheets']}
            value = 'GSM1234\n' * 5000
            tables['GEO Repository Links'] = [{'Publication Index': 'PUB-0001', 'Accession': 'GPL1', 'GEO Sample Accession': value}]
            output = Path(t)/'membership.xlsx'
            with self.assertRaises(ValueError):
                export_registered(template, output, tables)
            self.assertFalse(output.exists())
            receipt = export_registered(template, output, tables, overflow_sidecar=True)
            with gzip.open(receipt['overflow_sidecar'], 'rt', encoding='utf-8') as stream:
                self.assertEqual(json.load(stream)['records'][0]['value'], value)
            book = openpyxl.load_workbook(output)
            self.assertIn('record 1', book['GEO Repository Links']['D2'].value)
            self.assertEqual(book['GEO Repository Links']['C2'].value, 'GPL1')
            book.close()

    def test_complete_offline_synthesis_with_no_pmid_and_empty_repositories(self):
        with tempfile.TemporaryDirectory(dir=ROOT / 'outputs') as t:
            root = Path(t)
            (root / 'data/templates').mkdir(parents=True)
            for source in (ROOT / 'data/templates').glob('*-Enrichment.manifest.json'):
                shutil.copy2(source, root / 'data/templates' / source.name)
            for name in ('Enrichment-Synthesis.xlsx', 'Enrichment-Synthesis.manifest.json'):
                shutil.copy2(ROOT / 'data/templates' / name, root / 'data/templates' / name)
            pubs = [{'PUBLICATION_ID': '1', 'PMID': '123', 'TITLE': 'One', 'PUBLICATION_DATE': '2020-01-02'},
                    {'PUBLICATION_ID': '2', 'PMID': None, 'TITLE': 'Two', 'PUBLICATION_DATE': '2021'}]
            def save(name, value):
                path = root / name
                path.write_text(json.dumps(value), encoding='utf-8')
                return path
            index = save('index.json', pubs)
            registry = save('registry.json', publication_registry(pubs))
            pm = save('pm.json', {'rows': {'Publication': [['123', '', '', 'One', '2020-01-02']],
                'Repository Records': [], 'SRA Accessions': [], 'MeSH': [], 'MeSH Links': []}})
            attempts = [{'key': p.get('PMID') or 'ORCS:' + p['PUBLICATION_ID'],
                         'publication': p, 'observations': [], 'issues': [], 'status': 'complete'} for p in pubs]
            geo = save('geo.json', {'publications': {'123': {}, 'orcs-2': {}}, 'series': {}, 'samples': {}})
            ae = save('ae.json', {'publications': attempts, 'studies': {}, 'rows': {'Studies': []}})
            ena = save('ena.json', {'publications': attempts, 'studies': {}, 'issues': [], 'rows': {}})
            args = argparse.Namespace(root=root, stage='synthesize', publication_index=index, output_dir=root/'result',
                registry=registry, pubmed=pm, geo=geo, arrayexpress=ae, ena=ena, template=None)
            with patch('urllib.request.urlopen', side_effect=AssertionError('Synthesis attempted retrieval')):
                receipt = run_stage(args)
            self.assertEqual(receipt['sheet_rows']['Publications'], 2)
            book = openpyxl.load_workbook(root/'result/Enrichment-Synthesis.xlsx')
            self.assertEqual(len(book.sheetnames), 21)
            self.assertEqual(book['Publications']['A3'].value, 'PUB-0002')
            self.assertIsNone(book['Publications']['B3'].value)
            book.close()
            # The published offline export command must also work directly
            # from the compressed canonical output, without a pilot builder.
            from deathmap_ai.enrichment_export import main as export_main
            argv = ['deathmap-export-enrichment', '--database', str(root/'result/synthesis.json.gz'),
                    '--template', str(root/'data/templates/Enrichment-Synthesis.xlsx'),
                    '--output', str(root/'reexport.xlsx'), '--root', str(root)]
            with patch('sys.argv', argv), patch('urllib.request.urlopen', side_effect=AssertionError('Offline export used network')), redirect_stdout(io.StringIO()):
                export_main()
            self.assertTrue((root/'reexport.xlsx').exists())

    def test_pubmed_cache_reuse_keeps_native_retrieval_date(self):
        with tempfile.TemporaryDirectory(dir=ROOT / 'outputs') as t:
            root = Path(t)
            cache = root / 'old'
            cache.mkdir()
            body = b'<PubmedArticleSet/>'
            (cache/'one.xml').write_bytes(body)
            url = API + 'efetch.fcgi?db=pubmed'
            (cache/'requests.json').write_text(json.dumps([{'url': url, 'file': 'one.xml',
                'sha256': hashlib.sha256(body).hexdigest(), 'retrieved_utc': '2020-01-01T00:00:00Z'}]))
            transport = Retrieval(root/'new', [cache])
            with patch('deathmap_ai.pubmed_enrichment.urlopen', side_effect=AssertionError('Cache miss')):
                self.assertEqual(transport.request('efetch.fcgi', {'db': 'pubmed'}, 'one.xml'), body)
            self.assertEqual(transport.receipts[0]['retrieved_utc'], '2020-01-01T00:00:00Z')
            self.assertEqual(transport.attempts, 0)
            (cache/'one.xml').write_bytes(b'changed')
            with self.assertRaises(ValueError):
                Retrieval(root/'corrupt', [cache]).request('efetch.fcgi', {'db': 'pubmed'}, 'one.xml')

    def test_ena_full_title_fallback_stays_complete_and_bounded(self):
        from deathmap_ai.ena_collection import _recover_title, LimitReached
        class Fake:
            lock = threading.Lock()
            calls = []
            def get(self, url):
                self.calls.append(url)
                if 'portal' in url:
                    return b'error', {'file': 'error.response', 'url': url}
                return b'{"hitCount":0,"entries":[]}', {'file': 'empty.response', 'url': url}
        transport = Fake()
        title = 'A complete, punctuated publication title.'
        found, receipts = _recover_title(transport, title)
        self.assertEqual(found, [])
        from urllib.parse import urlsplit, parse_qs
        phrase = parse_qs(urlsplit(transport.calls[-1]).query)['query'][0]
        self.assertEqual(phrase, '"' + title + '"')
        transport.title_recoveries = 10
        with self.assertRaises(LimitReached):
            _recover_title(transport, title)


if __name__ == '__main__':
    unittest.main()
