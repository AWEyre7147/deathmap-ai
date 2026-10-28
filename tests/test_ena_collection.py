"""Inspect small native ENA fixtures for evidence and entity-boundary behavior."""
import hashlib
import json
import tempfile
import unittest
import openpyxl
from pathlib import Path
from deathmap_ai.ena_collection import Transport, _search, parse_study, _apply_metadata_only_policy
from deathmap_ai.ena_projection import project


class ENATests(unittest.TestCase):
    def test_registered_template_without_dimension_or_views_exports(self):
        from deathmap_ai.ena_workbook import export_workbook
        root = Path(__file__).resolve().parents[1]
        template = root / 'data/templates/ENA-Enrichment.xlsx'
        with tempfile.TemporaryDirectory() as temporary:
            folder = Path(temporary)
            registration = json.loads(template.with_suffix('.manifest.json').read_text(encoding='utf-8'))
            rows = {s['source_table']: [] for s in registration['sheets']}
            rows['Studies'] = [{'Study accession': 'PRJNA1', 'Study title': '=literal source', 'Runs in report': 0}]
            source = folder / 'database.json'
            source.write_text(json.dumps({'status': 'complete_with_omissions', 'rows': rows}), encoding='utf-8')
            output = folder / 'result.xlsx'
            result = export_workbook(source, template, output)
            book = openpyxl.load_workbook(output)
            self.assertEqual(book['Studies']['A2'].value, 'PRJNA1')
            # Bind assertions to native field names so approved column moves do
            # not turn the literal-text and zero-count checks into stale offsets.
            columns = next(s['columns'] for s in registration['sheets'] if s['name'] == 'Studies')
            positions = {c['source_key']: i + 1 for i, c in enumerate(columns)}
            title = book['Studies'].cell(2, positions['Study title'])
            self.assertEqual(title.value, '=literal source')
            self.assertEqual(title.data_type, 's')
            self.assertEqual(book['Studies'].cell(2, positions['Runs in report']).value, 0)
            self.assertEqual(result['sheet_rows']['Studies'], 1)
            book.close()

    def test_broad_project_policy_preserves_records_and_blank_counts(self):
        studies = {'PRJNA1': {'title': 'Broad collection', 'associations': [{'publication_key': '123'}], 'runs': [], 'run_count': 0},
                   'PRJNA2': {'title': 'Independent child', 'runs': [{'run_accession': 'SRR1'}], 'run_count': 1}}
        _apply_metadata_only_policy(studies, ['PRJNA1'])
        self.assertEqual(studies['PRJNA1']['associations'], [{'publication_key': '123'}])
        self.assertNotIn('run_count', studies['PRJNA1'])
        self.assertEqual(studies['PRJNA1']['run_report_status'], 'not-expanded-broad-collection')
        self.assertEqual(studies['PRJNA2']['run_count'], 1)

    def test_native_aliases_and_citation(self):
        record = parse_study('<PROJECT_SET><PROJECT accession="PRJNA1"><IDENTIFIERS><PRIMARY_ID>PRJNA1</PRIMARY_ID><SECONDARY_ID>SRP2</SECONDARY_ID><EXTERNAL_ID>GSE3</EXTERNAL_ID></IDENTIFIERS><TITLE>A</TITLE><PROJECT_LINK><XREF_LINK><DB>PUBMED</DB><ID>123</ID></XREF_LINK></PROJECT_LINK></PROJECT></PROJECT_SET>')
        self.assertEqual(record['aliases'], ['PRJNA1', 'SRP2'])
        self.assertEqual(record['pmids'], ['123'])
        self.assertEqual(record['identifiers'][-1]['value'], 'GSE3')

    def test_wrong_entity_and_ambiguous_record_rejected(self):
        for xml in ('<SET><SAMPLE/></SET>', '<SET><PROJECT/><PROJECT/></SET>'):
            with self.assertRaises(ValueError):
                parse_study(xml)

    def test_error_object_is_not_no_hits(self):
        class Fake:
            def get(self, url):
                return b'{"message":"Error in search"}', {}
        with self.assertRaises(ValueError):
            _search(Fake(), 'https://www.ebi.ac.uk/ena/portal/api/search', {}, 'complete-title')

    def test_hash_checked_cache_retains_retrieval_date(self):
        with tempfile.TemporaryDirectory() as temporary:
            root = Path(temporary)
            cache, run = root / 'cache', root / 'run'
            cache.mkdir(); run.mkdir()
            body = b'[]'
            (cache / 'one.response').write_bytes(body)
            (cache / 'requests.json').write_text(json.dumps([{'url': 'https://www.ebi.ac.uk/ena/portal/api/search?a=1&b=2', 'file': 'one.response', 'sha256': hashlib.sha256(body).hexdigest(), 'retrieved_utc': '2026-10-12'}]))
            transport = Transport(run, [cache], opener=lambda *args, **kwargs: self.fail('Network unexpectedly used'))
            actual, receipt = transport.get('https://www.ebi.ac.uk/ena/portal/api/search?b=2&a=1')
            self.assertEqual(actual, body)
            self.assertEqual(receipt['retrieved_utc'], '2026-10-12')
            (cache / 'one.response').write_bytes(b'changed')
            fresh = Transport(root, [cache], opener=lambda *args, **kwargs: self.fail('Network unexpectedly used'))
            # Existing run cache may legitimately supersede the corrupt older
            # source. A fresh output makes the tampered source the only cache.
            fresh.cache[next(iter(fresh.cache))] = (cache / 'one.response', {'sha256': hashlib.sha256(body).hexdigest()})
            with self.assertRaises(ValueError):
                fresh.get('https://www.ebi.ac.uk/ena/portal/api/search?a=1&b=2')

    def test_overlapping_runs_and_missing_file_metadata(self):
        row = {'run_accession': 'SRR1', 'study_accession': 'PRJNA2', 'secondary_study_accession': 'SRP2',
            'sample_accession': 'SAMN1', 'experiment_accession': 'SRX1', 'scientific_name': 'Homo sapiens',
            'fastq_ftp': 'example.fastq.gz', 'fastq_md5': '', 'fastq_bytes': '', 'read_count': '0', 'base_count': ''}
        def study(accession):
            return {'aliases': [accession], 'pmids': [], 'associations': [], 'title': 'A', 'description': 'B', 'organism': ['Homo sapiens'],
                'native_record_type': 'PROJECT', 'source_url': 'https://www.ebi.ac.uk/ena/browser/api/xml/' + accession,
                'hierarchy': [], 'runs': [dict(row)], 'run_report_response': 'raw.response'}
        database = {'studies': {'PRJNA1': study('PRJNA1'), 'PRJNA2': study('PRJNA2')}, 'publications': [], 'issues': []}
        rows = project(database)
        self.assertEqual(len(rows['Runs']), 1)
        self.assertEqual(rows['Runs'][0]['Reporting studies'], 'PRJNA1\nPRJNA2')
        self.assertEqual(rows['Runs'][0]['Read count'], 0)
        self.assertIsNone(rows['Runs'][0]['Base count'])
        self.assertIsNone(rows['File listings'][0]['MD5 checksum'])
        self.assertTrue(database['issues'])

    def test_distinct_native_values_are_not_merged(self):
        database = {'studies': {}, 'publications': [], 'issues': []}
        for i, title in enumerate(('First description', 'Different description'), 1):
            database['studies']['PRJNA' + str(i)] = {'aliases': [], 'pmids': [], 'associations': [], 'title': title, 'description': '',
                'organism': [], 'native_record_type': 'PROJECT', 'source_url': '', 'hierarchy': [],
                'runs': [{'run_accession': 'SRR1', 'study_accession': 'PRJNA1', 'secondary_study_accession': 'SRP1',
                         'sample_accession': 'SAMN1', 'experiment_accession': 'SRX1', 'sample_description': title}]}
        result = project(database)
        self.assertEqual(len(result['Runs']), 2)
        self.assertEqual(len(result['Samples']), 2)


if __name__ == '__main__':
    unittest.main()
