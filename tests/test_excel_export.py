"""Portable writer fixtures protect reviewer files, native links and template color."""
from copy import copy
import hashlib
import json
from pathlib import Path
import tempfile
import unittest
from openpyxl import load_workbook
from deathmap_ai.excel_export import export,ENTITY_SHEETS,body_style

ROOT=Path(__file__).resolve().parents[1]
TEMPLATE=ROOT/'data/templates/DeathMap-AI-output.xlsx'


class ExcelExportTests(unittest.TestCase):
    def projection(self):
        p={'sheets':{s:[] for s in ENTITY_SHEETS}}
        p['sheets']['Publications']=[{'publication_id':'PUB-0001','title_original':'=1+1','publication_year':2020,'publication_notes':None,'pmid_link':'https://pubmed.ncbi.nlm.nih.gov/123/','evidence_id_link':'EVI-00001'}]
        p['sheets']['Evidence']=[{'evidence_id':'EVI-00001','supports_entity_type':'publication','supports_entity_id':'PUB-0001','claim':'Fixture source claim','source_link':'https://example.org/source','evidence_status':'direct'}]
        return p

    def test_style_hyperlink_types_empty_entities_and_inputs_preserved(self):
        original_hash=hashlib.sha256(TEMPLATE.read_bytes()).hexdigest()
        with tempfile.TemporaryDirectory() as directory:
            run=Path(directory);(run/'projection.json').write_text(json.dumps(self.projection()))
            result=export(run,TEMPLATE);source=load_workbook(TEMPLATE);book=load_workbook(result)
            self.assertEqual(source.sheetnames,book.sheetnames)
            for name in ENTITY_SHEETS:
                for a,b in zip(source[name][1],book[name][1]):
                    for attr in ['font','fill','border','alignment','number_format']:self.assertEqual(copy(getattr(a,attr)),copy(getattr(b,attr)))
            self.assertEqual(book['Publications']['B2'].value,'=1+1');self.assertEqual(book['Publications']['B2'].data_type,'s')
            self.assertEqual(book['Publications']['E2'].value,2020)
            self.assertEqual(book['Evidence']['G2'].hyperlink.target,'https://example.org/source')
            self.assertEqual(book['Evidence']['G2'].data_type,'s')
            self.assertIsNone(book['Publications']['O2'].value)
            for name in ['Screen Groups','Screens','Datasets','Screen-Dataset Links']:
                self.assertFalse(any(c.value is not None for row in book[name].iter_rows(min_row=2) for c in row))
            self.assertEqual(body_style(source['Publications'],2),book['Publications']['B2']._style)
            self.assertEqual(source.loaded_theme,book.loaded_theme)
            source.close();book.close()
        self.assertEqual(hashlib.sha256(TEMPLATE.read_bytes()).hexdigest(),original_hash)

    def test_existing_owner_workbook_refused_before_projection_or_template_read(self):
        with tempfile.TemporaryDirectory() as directory:
            run=Path(directory);path=run/'workbook-authored.xlsx';path.write_bytes(b'owner reviewed bytes')
            with self.assertRaises(FileExistsError):export(run,run/'missing-template.xlsx')
            self.assertEqual(path.read_bytes(),b'owner reviewed bytes')

    def test_deleted_body_rows_retain_grouped_column_colors(self):
        from openpyxl import Workbook
        from openpyxl.styles import PatternFill
        book=Workbook();sheet=book.active
        sheet.column_dimensions.group('B','D')
        sheet.column_dimensions['B'].fill=PatternFill('solid',fgColor='FFFF00')
        self.assertEqual(body_style(sheet,3),sheet.column_dimensions['B']._style)
        sheet['C2'].fill=PatternFill('solid',fgColor='F4B183')
        self.assertEqual(body_style(sheet,3),sheet['C2']._style)
        book.close()

    def test_generated_review_notes_and_unmapped_fields_fail_without_output(self):
        for field in ['publication_notes','unknown_scientific_field']:
            with tempfile.TemporaryDirectory() as directory:
                run=Path(directory);p=self.projection();p['sheets']['Publications'][0][field]='unexpected value'
                (run/'projection.json').write_text(json.dumps(p))
                with self.assertRaises(ValueError):export(run,TEMPLATE)
                self.assertFalse((run/'workbook-authored.xlsx').exists())
