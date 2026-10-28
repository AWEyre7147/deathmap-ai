"""Small template-population checks; no retrieval or database regeneration."""
from copy import copy
import json
from pathlib import Path
import tempfile
import unittest
import openpyxl
from deathmap_ai.geo_workbook import export_workbook

ROOT = Path(__file__).resolve().parents[1]
TEMPLATE = ROOT / 'data/templates/GEO-Enrichment.xlsx'
MANIFEST = TEMPLATE.with_suffix('.manifest.json')


class GeoWorkbookTests(unittest.TestCase):
    def setUp(self):
        # Use a self-cleaning workspace without recreating retired GEO run folders.
        self.directory = tempfile.TemporaryDirectory(dir=ROOT / 'outputs')
        self.addCleanup(self.directory.cleanup)
        self.output = Path(self.directory.name) / 'fixture.xlsx'
        self.tables = {
            'Series Records': [{'PMID': '1', 'geo_accession': 'GSE2', 'sample_count': 1}],
            'Series Hierarchy': [{'PMID': '1', 'GSE': 'GSE2', 'Relationship': 'SubSeries of', 'Related GSE': 'GSE3'}],
            'Repository Identifiers': [{'PMID': '1', 'GSE': 'GSE2', 'GSM': 'GSM4'}],
            'Sample Metadata': [{'PMID': '1', 'GSE': 'GSE2', 'GSM': 'GSM4', 'scan_protocol': '=literal metadata'}],
        }

    def test_population_preserves_owner_format_and_last_field(self):
        before = TEMPLATE.read_bytes()
        receipt = export_workbook(TEMPLATE, self.output, self.tables)
        book = openpyxl.load_workbook(self.output)
        template = openpyxl.load_workbook(TEMPLATE)
        try:
            self.assertIn(book['Series Records']['L2'].value, (None, ''))
            headers = [c.value for c in book['Sample Metadata'][1]]
            scan = book['Sample Metadata'].cell(2, headers.index('Scan Protocol') + 1)
            self.assertEqual(scan.value, '=literal metadata')
            self.assertEqual(scan.data_type, 's')
            self.assertEqual(copy(book['Sample Metadata']['A2'].font), copy(template['Sample Metadata']['A2'].font))
            self.assertIsNone(book['Sample Metadata'].row_dimensions[2].height)
            self.assertEqual(receipt['sheet_rows']['Sample Metadata'], 1)
        finally:
            book.close()
            template.close()
        self.assertEqual(TEMPLATE.read_bytes(), before)
        with self.assertRaises(FileExistsError):
            export_workbook(TEMPLATE, self.output, self.tables)

    def test_changed_registration_hash_is_rejected(self):
        registration = json.loads(MANIFEST.read_text(encoding='utf-8'))
        registration['sha256'] = '0' * 64
        manifest = Path(self.directory.name) / 'wrong.json'
        manifest.write_text(json.dumps(registration), encoding='utf-8')
        with self.assertRaisesRegex(ValueError, 'register'):
            export_workbook(TEMPLATE, self.output, self.tables, manifest)
        self.assertFalse(self.output.exists())

    def test_missing_table_is_not_a_successful_empty_sheet(self):
        with self.assertRaisesRegex(ValueError, 'Missing projected table'):
            export_workbook(TEMPLATE, self.output, {})

    def test_oversize_text_stops_before_export_without_truncation(self):
        self.tables['Sample Metadata'][0]['scan_protocol'] = 'x' * 32768
        with self.assertRaisesRegex(ValueError, 'cell limit'):
            export_workbook(TEMPLATE, self.output, self.tables)
        self.assertFalse(self.output.exists())


if __name__ == '__main__':
    unittest.main()
