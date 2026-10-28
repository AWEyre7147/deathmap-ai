"""Display-label and registration checks preserve source-field identities."""
from pathlib import Path
import hashlib,json,tempfile,unittest
from openpyxl import load_workbook
from deathmap_ai.display_labels import readable_title,readable_headers,merge_display_values,identifier_label
from deathmap_ai.enrichment_export import export_registered

ROOT=Path(__file__).resolve().parents[1]

class DisplayLabelTests(unittest.TestCase):
    def test_acronyms_channel_suffix_and_scientific_qualifiers(self):
        self.assertEqual(readable_title('source_name_ch1'),'Source Name')
        self.assertEqual(readable_title('library_scope_normalized'),'Library Scope Normalized')
        self.assertEqual(readable_title('RNA interference'),'RNA Interference')
        self.assertEqual(readable_title('hyb_protocol'),'Hybridization Protocol')
        self.assertEqual(readable_title('Size (bytes)'),'Size (Bytes)')

    def test_case_variants_have_the_same_normalized_label(self):
        self.assertEqual(readable_headers(['Age','age','RNA interference','Rna interference']),['Age','Age','RNA Interference','RNA Interference'])
        self.assertEqual(readable_headers(['Age','Age (2)']),['Age','Age'])
        self.assertEqual(readable_headers([None,None,'pmid']),[None,None,'PMID'])

    def test_consolidation_preserves_conflicts_and_deduplicates_matches(self):
        self.assertEqual(merge_display_values(['adult', 'adult', None]), 'adult')
        self.assertEqual(merge_display_values(['adult', '10']), 'adult\n10')
        self.assertEqual(merge_display_values([0, None]), 0)

    def test_arrayexpress_duplicate_field_export(self):
        template=ROOT/'data/templates/BioStudies-ArrayExpress-Enrichment.xlsx'
        manifest=json.loads(template.with_suffix('.manifest.json').read_text())
        tables={s['source_table']:[] for s in manifest['sheets']}
        tables['Sample metadata']=[{'Study accession':'S1','Age':'adult','age':'10'}]
        with tempfile.TemporaryDirectory(dir=ROOT/'outputs') as folder:
            out=Path(folder)/'export.xlsx'
            export_registered(template,out,tables)
            w=load_workbook(out);s=w['Sample metadata'];headers=[c.value for c in s[1]]
            self.assertEqual(headers.count('Age'),1)
            self.assertNotIn('Alternate Age',headers)
            self.assertEqual(s.cell(2,headers.index('Age')+1).value,'adult\n10')
            w.close()

    def test_identifier_labels_follow_namespace_not_provider(self):
        self.assertEqual(identifier_label('GEO','Series Records','geo_accession',''), 'GEO Series Accession')
        self.assertEqual(identifier_label('GEO','Sample Metadata','geo_accession',''), 'GEO Sample Accession')
        self.assertEqual(identifier_label('ENA','Studies','Study accession',''), 'BioProject Accession')
        self.assertEqual(identifier_label('ENA','Runs','SRA study accession',''), 'SRA Study Accession')
        self.assertEqual(identifier_label('ENA','Samples','Sample accession',''), 'BioSample Accession')
        self.assertEqual(identifier_label('PubMed','SRA Accessions','Sample Accession',''), 'SRA Sample Accession')

    def test_synthesis_projection_matches_all_registered_display_keys(self):
        from deathmap_ai.enrichment_commands import synthesis_tables
        registration=json.loads((ROOT/'data/templates/Enrichment-Synthesis.manifest.json').read_text())
        data={'tables':{}}
        for spec in registration['sheets']:
            row={}
            for column in spec['columns']:
                for key in column.get('native_source_keys',[column['native_source_key']]):
                    row[key]='value:'+key
            data['tables'][spec['native_source_table']]=[row]
        projected=synthesis_tables(data,ROOT)
        for spec in registration['sheets']:
            self.assertEqual(list(projected[spec['name']][0]),[c['source_key'] for c in spec['columns']],spec['name'])
            native = data['tables'][spec['native_source_table']][0]
            for c in spec['columns']:
                keys = c.get('native_source_keys', [c['native_source_key']])
                expected = merge_display_values(native[k] for k in keys)
                self.assertEqual(projected[spec['name']][0][c['source_key']], expected, (spec['name'],c['header']))

    def test_registered_templates_have_matching_hashes_and_headers(self):
        for path in (ROOT/'data/templates').glob('*.manifest.json'):
            manifest=json.loads(path.read_text(encoding='utf-8'))
            self.assertNotIn('status',manifest)
            if 'template' not in manifest or 'sheets' not in manifest:continue
            self.assertNotIn(chr(92), manifest['template'], 'Template paths must use portable forward slashes')
            template=ROOT/manifest['template']
            self.assertEqual(hashlib.sha256(template.read_bytes()).hexdigest(),manifest['sha256'])
            w=load_workbook(template,read_only=True)
            for spec in manifest['sheets']:
                headers=[c.value for c in w[spec['name']][1]]
                self.assertEqual(headers,spec['headers'])
                self.assertEqual(headers,readable_headers(headers))
            w.close()

    def test_geo_export_keeps_native_channel_keys(self):
        template=ROOT/'data/templates/GEO-Enrichment.xlsx'
        manifest=json.loads(template.with_suffix('.manifest.json').read_text())
        tables={s['source_table']:[] for s in manifest['sheets']}
        tables['Sample Metadata']=[{'PMID':'123','GSM':'GSM1','source_name_ch1':'=literal source','organism_ch1':'human'}]
        with tempfile.TemporaryDirectory(dir=ROOT/'outputs') as folder:
            out=Path(folder)/'export.xlsx'
            export_registered(template,out,tables)
            w=load_workbook(out);s=w['Sample Metadata'];headers=[c.value for c in s[1]]
            self.assertEqual(s.cell(2,headers.index('Source Name')+1).value,'=literal source')
            self.assertEqual(s.cell(2,headers.index('Organism')+1).value,'human')
            self.assertEqual(s.cell(2,headers.index('Source Name')+1).data_type,'s')
            w.close()

if __name__=='__main__':unittest.main()
