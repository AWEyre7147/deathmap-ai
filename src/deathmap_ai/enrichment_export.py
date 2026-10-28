"""Stream registered enrichment tables into a new, verified Excel workbook.

The template's native XML envelope and headers survive unchanged. Body cells
are written incrementally as literal text, numbers or Booleans, avoiding a full
spreadsheet-engine model and its timestamp coercion. No retrieval occurs here.
"""
import argparse
import copy
import hashlib
import gzip
import json
from pathlib import Path
import re
import zipfile
import xml.etree.ElementTree as ET
from xml.sax.saxutils import escape

from .display_labels import merge_display_values

import openpyxl
from openpyxl.utils import get_column_letter
from .pubmed_workbook import NS, q, replace_block


def export_registered(template, output, tables, *, overflow_sidecar=False):
    """Write mapped tables into a new registered template and verify every cell.

    Parameters
    ----------
    template, output : Path or str
        Registered header-only input and nonexisting workbook destination.
    tables : dict
        Table names mapped to record dictionaries or positional rows. The
        registration supplies source_key or source_index for each column.
    overflow_sidecar : bool, optional
        Preserve oversized values in a lossless compressed sidecar and place
        explicit references in cells. Default False stops before writing.

    Returns
    -------
    dict
        Template/output hashes, row counts and independently verified cell count.

    Raises
    ------
    ValueError
        Registration, mapping, cell length, native type or saved value is invalid.
    FileExistsError
        Output already exists; no owner workbook is overwritten.
    """
    template, output = Path(template), Path(output)
    manifest = json.loads(template.with_suffix('.manifest.json').read_text(encoding='utf-8'))
    digest = hashlib.sha256(template.read_bytes()).hexdigest()
    if digest != manifest['sha256']:
        raise ValueError('Registered template hash mismatch')
    if output.exists():
        raise FileExistsError(output)
    overflows, replacements = [], {}
    def rows(spec):
        name = spec['source_table'].removeprefix('rows.')
        if name not in tables:
            raise ValueError('Missing projected table: ' + name)
        for row_index, record in enumerate(tables[name]):
            yield [replacements.get((spec['name'], row_index, i),
                   merge_display_values(record.get(key) for key in c.get('source_keys', [c.get('source_key', c['header'])])) if isinstance(record, dict)
                   else record[c.get('source_column_index', c.get('source_index', i))]) for i, c in enumerate(spec['columns'])]
    # Preflight before creating output: no silent truncation, missing mappings,
    # experimental files or executable formula interpretation.
    counts = {}
    for spec in manifest['sheets']:
        counts[spec['name']] = 0
        for row_index, row in enumerate(rows(spec)):
            for column, value in enumerate(row):
                if not isinstance(value, (str, int, float, bool, type(None))):
                    raise ValueError('Unsupported cell type')
                if isinstance(value, str) and re.search(r'[\x00-\x08\x0b\x0c\x0e-\x1f]', value):
                    raise ValueError('Native text contains invalid XML controls; evidence retained')
                if isinstance(value, str) and len(value) > 32767:
                    if not overflow_sidecar:
                        raise ValueError('Native text exceeds Excel cell constraints; evidence retained')
                    number = len(overflows) + 1
                    overflows.append({'record': number, 'sheet': spec['name'], 'excel_row': row_index + 2,
                        'column': spec['headers'][column], 'source_table': spec['source_table'],
                        'source_key': spec['columns'][column].get('source_key', spec['headers'][column]),
                        'value': value, 'characters': len(value)})
                    replacements[(spec['name'], row_index, column)] = ('Full value exceeds Excel cell limit; see ' +
                        output.with_suffix('.overflow.json.gz').name + ', record ' + str(number))
            counts[spec['name']] += 1
    output.parent.mkdir(parents=True, exist_ok=True)
    if overflows:
        sidecar = output.with_suffix('.overflow.json.gz')
        if sidecar.exists():
            raise FileExistsError(sidecar)
        with gzip.open(sidecar, 'wt', encoding='utf-8') as stream:
            json.dump({'schema_version': 'excel-native-overflow/1', 'records': overflows}, stream, ensure_ascii=False)
    with zipfile.ZipFile(template) as source:
        styles = ET.fromstring(source.read('xl/styles.xml'))
        xfs = styles.find(q('cellXfs'))
        body_style = len(xfs)
        xfs.append(copy.deepcopy(xfs[0]))
        xfs.set('count', str(len(xfs)))
        style_bytes = replace_block(source.read('xl/styles.xml').decode(), 'cellXfs', xfs).encode()
        workbook = ET.fromstring(source.read('xl/workbook.xml'))
        rels = {r.get('Id'): r.get('Target') for r in ET.fromstring(source.read('xl/_rels/workbook.xml.rels'))}
        sheets = list(workbook.find(q('sheets')))
        if [s.get('name') for s in sheets] != [s['name'] for s in manifest['sheets']]:
            raise ValueError('Registered sheet order differs')
        members = {}
        for element, spec in zip(sheets, manifest['sheets']):
            target = rels[element.get('{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id')]
            member = target.lstrip('/') if target.startswith('/') else 'xl/' + target
            members[member] = spec
        with output.open('xb') as stream, zipfile.ZipFile(stream, 'w', zipfile.ZIP_DEFLATED) as result:
            for item in source.infolist():
                if item.filename == 'xl/styles.xml':
                    result.writestr(item, style_bytes)
                elif item.filename not in members:
                    result.writestr(item, source.read(item.filename))
                else:
                    spec = members[item.filename]
                    xml = source.read(item.filename).decode()
                    node = ET.fromstring(xml)
                    data = node.find(q('sheetData'))
                    if len(data) != 1:
                        raise ValueError('Expected header-only template')
                    # Keep the original header bytes and worksheet properties;
                    # only the data extent and appended body rows change.
                    block = re.search(r'<(?:\w+:)?sheetData\b[^>]*>(.*?)</(?:\w+:)?sheetData>', xml, re.S)
                    header = block[1]
                    extent = 'A1:' + get_column_letter(len(spec['headers'])) + str(counts[spec['name']] + 1)
                    dimension = '<dimension xmlns="' + NS + '" ref="' + extent + '"/>'
                    xml, replaced = re.subn(r'<(?:\w+:)?dimension\b[^>]*/>', dimension, xml, count=1)
                    if not replaced:
                        xml = re.sub(r'(<(?:\w+:)?(?:sheetViews|sheetFormatPr|cols|sheetData)\b)', lambda m: dimension + m[1], xml, count=1)
                    block = re.search(r'<(?:\w+:)?sheetData\b[^>]*>(.*?)</(?:\w+:)?sheetData>', xml, re.S)
                    with result.open(item.filename, 'w') as body:
                        body.write(xml[:block.start(1)].encode())
                        body.write(header.encode())
                        for number, row in enumerate(rows(spec), 2):
                            cells = []
                            for column, value in enumerate(row, 1):
                                address = get_column_letter(column) + str(number)
                                start = f'<c r="{address}" s="{body_style}"'
                                if value is None or value == '':
                                    cell = start + '/>'
                                elif isinstance(value, bool):
                                    cell = start + ' t="b"><v>' + str(int(value)) + '</v></c>'
                                elif isinstance(value, (int, float)):
                                    cell = start + ' t="n"><v>' + str(value) + '</v></c>'
                                else:
                                    cell = start + ' t="inlineStr"><is><t xml:space="preserve">' + escape(value) + '</t></is></c>'
                                cells.append(cell)
                            body.write((f'<row xmlns="{NS}" r="{number}">' + ''.join(cells) + '</row>').encode())
                        body.write(xml[block.end(1):].encode())
    verified = 0
    book = openpyxl.load_workbook(output, read_only=True)
    try:
        for spec in manifest['sheets']:
            iterator = book[spec['name']].iter_rows(values_only=True)
            if list(next(iterator)) != spec['headers']:
                raise ValueError('Saved header mismatch')
            for expected in rows(spec):
                actual = next(iterator)
                if ['' if v is None else v for v in actual] != ['' if v is None else v for v in expected]:
                    raise ValueError('Saved value mismatch: ' + spec['name'])
                verified += len(expected)
            if next(iterator, None) is not None:
                raise ValueError('Unexpected saved rows')
    finally:
        book.close()
    receipt = {'template': str(template), 'template_sha256': digest, 'output': str(output),
        'workbook_sha256': hashlib.sha256(output.read_bytes()).hexdigest(), 'sheet_rows': counts,
        'verified_cells': verified, 'writer': 'streamed-template-XML/1'}
    if overflows:
        receipt.update(overflow_sidecar=str(sidecar), overflow_values=len(overflows),
                       overflow_sha256=hashlib.sha256(sidecar.read_bytes()).hexdigest())
    output.with_suffix('.manifest.json').write_text(json.dumps(receipt, indent=2), encoding='utf-8')
    return receipt


def main():
    """Export saved canonical rows or a synthesis payload without retrieval."""
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--database', type=Path, required=True)
    parser.add_argument('--template', type=Path, required=True)
    parser.add_argument('--output', type=Path, required=True)
    parser.add_argument('--root', type=Path, default=Path(__file__).resolve().parents[2])
    args = parser.parse_args()
    if args.database.suffix == '.gz':
        with gzip.open(args.database, 'rt', encoding='utf-8') as stream:
            data = json.load(stream)
    else:
        data = json.loads(args.database.read_text(encoding='utf-8'))
    if data.get('schema_version', '').startswith('enrichment-synthesis/'):
        from .enrichment_commands import synthesis_tables
        tables = synthesis_tables(data, args.root)
    else:
        tables = data.get('rows', data.get('tables'))
    receipt = export_registered(args.template, args.output, tables,
        overflow_sidecar=data.get('schema_version', '').startswith('enrichment-synthesis/'))
    receipt.update(source_database=str(args.database.resolve()),
                   source_database_sha256=hashlib.sha256(args.database.read_bytes()).hexdigest())
    args.output.with_suffix('.manifest.json').write_text(json.dumps(receipt, indent=2), encoding='utf-8')
    print(json.dumps(receipt))


if __name__ == '__main__':
    main()
