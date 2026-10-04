"""Verify the archived ORCS workbook before publishing it to the owner's path.

Checks resolve every projected identity/foreign key, preserve existing values and
formulas, and rehash the local sources and prior resource outputs. This command
performs no retrieval and does not modify the workbook.
"""
from pathlib import Path
import argparse
import json
from datetime import datetime,timezone
from openpyxl import load_workbook
from deathmap_ai.orcs_filters import read,sha,write
from deathmap_ai.orcs_paths import resolve_historical


def verify(run, root):
    """Return verification counts or fail before replacing the owner's workbook."""
    run,root=Path(run),Path(root);manifest=read(run/'run_manifest.json');projection=read(run/'projection.json')
    if projection.get('version','').startswith('orcs-publication-enrichment-v1.'):
        from .orcs_publication_enrichment import verify_enriched
        return verify_enriched(run, root)
    original=load_workbook(run/'workbook-original.xlsx',data_only=False)
    output=load_workbook(run/'DeathMap-AI-v1-reference-output.xlsx',data_only=False)
    if original.sheetnames!=output.sheetnames: raise ValueError('Sheet names/order changed')
    counts={};formula_count=0
    for before in original:
        after=output[before.title]
        for row in before:
            for cell in row:
                if cell.value is not None and after[cell.coordinate].value!=cell.value:
                    raise ValueError(f'Existing cell changed: {before.title}!{cell.coordinate}')
                if cell.data_type=='f':formula_count+=1
                if cell.has_style:
                    other=after[cell.coordinate]
                    if cell.font!=other.font or cell.border!=other.border or cell.number_format!=other.number_format:
                        raise ValueError(f'Existing cell style changed: {before.title}!{cell.coordinate}')
        headers=[c.value for c in after[1]];actual=[dict(zip(headers,values)) for values in list(after.values)[1:] if any(v is not None for v in values)]
        expected=projection['sheets'].get(before.title)
        if expected is None:continue
        key=projection['headers'][before.title][0]
        byid={x[key]:x for x in actual}
        if len(byid)!=len(actual) or len(actual)!=len(expected):raise ValueError('Row reconciliation failed: '+before.title)
        for row in expected:
            for field,value in row.items():
                observed=byid[row[key]].get(field)
                if field=='retrieved_at' and isinstance(observed,datetime) and value:
                    expected_date=datetime.fromisoformat(value).replace(tzinfo=None)
                    if abs((observed-expected_date).total_seconds()) < 0.002:continue
                # Null remains empty, while identifiers and reported text remain exact.
                if observed!=value and not (value in (None,'') and observed in (None,'')):
                    raise ValueError(f'Projection mismatch: {before.title} / {row[key]} / {field}')
        counts[before.title]=len(actual)
    screens={row[0] for row in list(output['Screens'].values)[1:]};pubs={row[0] for row in list(output['Publications'].values)[1:]}
    for row in projection['sheets']['Screens']:
        if row.get('publication_id') and row['publication_id'] not in pubs:raise ValueError('Screen publication missing')
    for row in projection['sheets']['Sources & Evidence']:
        if row.get('supports_entity_type')=='screen' and row['supports_entity_id'] not in screens:raise ValueError('Evidence screen missing')
    for row in projection['sheets']['Publications']:
        if row.get('pmid_link') and row['pmid_link']!='https://pubmed.ncbi.nlm.nih.gov/'+str(row['pmid'])+'/':raise ValueError('PMID URL mismatch')
    review=output['Screens'];headers=[x.value for x in review[1]]
    if headers[-2:]!=['reviewer_decision','reviewer_notes']:raise ValueError('Reviewer fields not at far right')
    for row in list(review.rows)[1:]:
        if row[0].value in projection['roles'] and row[-2].value in (None,''):
            if not str(row[-1].fill.fgColor.rgb).endswith('F4B183'):raise ValueError('Requested reviewer cell not orange')
    baseline=read(run/'preservation.json')['before'];changed=[p for p,h in baseline.items() if not resolve_historical(root/p).is_file() or sha(root/p)!=h]
    if changed:raise ValueError('Historical/source files changed: '+str(changed[:8]))
    for name,h in manifest['input_hashes'].items():
        if name.endswith('.xlsx'):
            if sha(run/'workbook-original.xlsx')!=h:raise ValueError('Workbook backup mismatch')
        elif name in ('searches/orcs/cancer-cell-crispr-knockout-v01.json', 'searches/orcs/cancer-cell-crispr-knockout-v01/profile.json'):
            # A historical run owns its exact profile snapshot; the active
            # profile may subsequently relocate an unchanged metadata input.
            if sha(run/'profile.json')!=h:raise ValueError('Historical profile snapshot changed')
        elif sha(root/name)!=h:raise ValueError('Input changed: '+name)
    result={'verified_at':datetime.now(timezone.utc).isoformat(),'row_counts':counts,'existing_formulas_preserved':formula_count,
      'historical_and_reference_files_unchanged':len(baseline),'original_workbook_sha256':sha(run/'workbook-original.xlsx'),
      'populated_workbook_sha256':sha(run/'DeathMap-AI-v1-reference-output.xlsx'),'network_requests':0,
      'reviewer_fields_verified':True,'projection_reconciled':True}
    write(run/'verification.json',result);return result


def main():
    """Verify a run without publishing or changing any original files."""
    p=argparse.ArgumentParser();p.add_argument('run',type=Path);p.add_argument('--root',type=Path,default=Path(__file__).resolve().parents[2]);a=p.parse_args();print(json.dumps(verify(a.run,a.root),indent=2))
if __name__=='__main__':main()
