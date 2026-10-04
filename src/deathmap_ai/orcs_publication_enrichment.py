"""Project saved ORCS bibliography into a new copy of a completed filter workbook.

Native screen associations own the join; titles never identify entities. This
offline stage preserves filter roles, owner values and source dates, and creates
publication/association evidence without manufacturing experimental datasets.
"""
import argparse
import copy
import json
import re
import shutil
from datetime import date, datetime, timezone
from pathlib import Path
from urllib.parse import quote

from openpyxl import load_workbook
from .orcs_filter_projection import existing_rows, stable_id, DIRECT, CROSSWALKS
from .orcs_filters import read, sha, write

VERSION = 'orcs-publication-enrichment-v1.0'
REVIEW = ['reviewer_decision', 'reviewer_notes']
FINGERPRINTS = {
    'data/orcs/publication-index.json': '0ebd60824abe9832d5eb45f3e65a8765c4a864af13c999476eb4deecb194e76d',
    'data/orcs/publication-screen-links.json': 'cf6ad9bad3d657781969e77b4d9ff5b283a8da8bc05a4721477db7d5ce2aac2b',
    'docs/orcs/workbook-field-source-map.md': 'c18ba5b8b443c96327a4d4213bd414eb577ac50d708f499d8fba22c9c3406f08',
    'data/templates/DeathMap-AI-v1-reference-output.xlsx': 'c8e40370121c8ac2a5d7d4480f2dc8ac31c6153a46d473ac8682939762f025ca',
}


def unique(rows, key):
    """Index required native identities, rejecting missing and duplicate records."""
    result = {}
    for row in rows:
        identity = row.get(key)
        if not isinstance(identity, str) or not identity or identity in result:
            raise ValueError('Missing/duplicate identity: ' + key)
        result[identity] = row
    return result


def validate_catalogs(screens, publications, links):
    """Validate exact source identities and foreign keys across all local catalogs."""
    si, pi, li = unique(screens, 'SCREEN_ID'), unique(publications, 'PUBLICATION_ID'), unique(links, 'SCREEN_ID')
    if set(si) != set(li):
        raise ValueError('Missing/extra screen associations')
    for sid, link in li.items():
        if link['PUBLICATION_ID'] not in pi:
            raise ValueError('Association publication foreign key missing')
        pub = pi[link['PUBLICATION_ID']]
        for field in ['SOURCE_TYPE', 'SOURCE_ID']:
            if not link.get(field) or link[field] != si[sid].get(field) or link[field] != pub.get(field):
                raise ValueError(f'Source identity disagreement: {sid}/{field}')
        if sid not in pub['SCREEN_IDS']:
            raise ValueError('Publication screen membership disagreement')
    for pid, pub in pi.items():
        if set(pub['SCREEN_IDS']) != {sid for sid, link in li.items() if link['PUBLICATION_ID'] == pid}:
            raise ValueError('Publication membership mismatch: ' + pid)
        if pub.get('PMID') and not (pub['SOURCE_TYPE'] == pub['PAGE_SOURCE_TYPE'] == 'pubmed'
                                   and pub['SOURCE_ID'] == pub['PAGE_SOURCE_ID'] == pub['PMID']):
            raise ValueError('Unsupported PMID promotion: ' + pid)
    return si, pi, li


def field_map(text):
    """Keep all exact mapping headers and classify their scientific support."""
    result = []; sheet = None
    for line in text.splitlines():
        if line.startswith('### '): sheet = line[4:]
        if line.startswith('| ') and sheet:
            parts = [x.strip() for x in line.strip('|').split('|')]
            if len(parts) != 3 or parts[0] == 'Exact workbook header': continue
            status = parts[2]
            classification = ('review-dependent' if 'Review' in status or 'review' in status else
                              'unsupported' if status in ('Unresolved', 'Asset metadata only', 'Run-dependent') else
                              'provenance' if 'Provenance' in status or 'provenance' in status else
                              'explicitly derived' if 'Derived' in status or 'derived' in status else 'reported')
            result.append({'sheet': sheet, 'field': parts[0], 'classification': classification,
                           'source_rule': parts[1], 'source_status': status})
    return result


def enrich(original, workbook, screens, publications, links, projection_time, association_retrieved_at=None):
    """Return a canonical projection and cell audit without writing or retrieving.

    `original` is the frozen projection; `workbook` contains current owner cells.
    Required native catalogs must already have passed provenance/hash validation.
    Nonempty cells win; safe note/evidence appends retain their exact old prefix.
    """
    si, pi, li = validate_catalogs(screens, publications, links)
    p = copy.deepcopy(original)
    p['version'] = VERSION
    p['sheets'] = {name: copy.deepcopy(value['rows']) for name, value in workbook.items()}
    # Excel stores dates at millisecond precision. Preserve the full original
    # receipt timestamp in JSON when it agrees with the copied workbook date.
    old_evidence = {row['evidence_id']:row for row in original['sheets']['Sources & Evidence']}
    for row in p['sheets']['Sources & Evidence']:
        value = row.get('retrieved_at')
        if isinstance(value, datetime):
            saved = old_evidence.get(row['evidence_id'],{}).get('retrieved_at')
            if saved and abs((value-datetime.fromisoformat(saved).replace(tzinfo=None)).total_seconds()) < .002:
                row['retrieved_at'] = saved
            else: row['retrieved_at'] = value.isoformat()
    p['headers'] = {name: value['headers'][:] for name, value in workbook.items()}
    p['cell_audit'] = []; p['approved_appends'] = []; p['publication_review'] = {}
    p['existing_value_conflicts'] = copy.deepcopy(original.get('existing_value_conflicts', []))
    p['source_projection_issues'] = copy.deepcopy(original.get('issues', [])); p['issues'] = []
    evidence = p['sheets']['Sources & Evidence']
    pub_rows = {row['publication_id']: row for row in p['sheets']['Publications']}
    screen_rows = {row['screen_id']: row for row in p['sheets']['Screens']}
    if set(screen_rows) != set(original['roles']): raise ValueError('Workbook screen set differs from frozen projection')
    mapping = []; used = {}; evidence_ids = {row['evidence_id'] for row in evidence}

    def propose(sheet, row, field, value, locator, transformation='reported', append=False):
        key = p['headers'][sheet][0]; retained = row.get(field)
        audit = {'sheet': sheet, 'id': row[key], 'field': field, 'retained': retained,
                 'proposed': value, 'source_locator': locator, 'transformation': transformation}
        if value in (None, ''):
            audit['status'] = 'unresolved'; audit['reason'] = 'No supported value in accepted local inputs'
        elif retained in (None, ''):
            row[field] = value; audit['status'] = 'added'
        elif retained == value:
            audit['status'] = 'already supported'
        elif append and isinstance(retained, str) and not retained.startswith('='):
            if value in retained: audit['status'] = 'already supported'
            else:
                row[field] = retained + '\n' + str(value); audit['status'] = 'appended'
                p['approved_appends'].append({**audit, 'value': row[field]})
        else:
            audit['status'] = 'retained conflict'
            p['existing_value_conflicts'].append({**audit, 'conflict_status': 'owner value retained; source proposal available'})
        p['cell_audit'].append(audit)

    def add_evidence(kind, native, entity_type, entity_id, locator, text, retrieved, status):
        eid = stable_id('EVI', VERSION, kind, native)
        if eid in evidence_ids:
            existing = next(r for r in evidence if r['evidence_id']==eid)
            if existing['supports_entity_type']!=entity_type or existing['supports_entity_id']!=entity_id:
                raise ValueError('Evidence identity collision')
            return eid
        evidence_ids.add(eid)
        evidence.append({'evidence_id': eid, 'source_type_normalized': 'discovery_resource',
            'source_name': 'BioGRID ORCS saved publication metadata' if kind == 'header' else 'BioGRID ORCS explicit association',
            'native_record_id': native, 'retrieved_at': retrieved,
            'retrieval_method': 'offline local reuse; zero new network requests', 'tool_or_package': VERSION,
            'source_locator': locator, 'evidence_text_summary': json.dumps(text, ensure_ascii=False),
            'supports_entity_type': entity_type, 'supports_entity_id': entity_id, 'evidence_status': status,
            'evidence_notes': 'Offline projection at ' + projection_time + '; source retrieval timestamp retained'})
        return eid

    # Existing screen foreign keys and frozen identifier rows jointly anchor the
    # 92 historical IDs. An owner-edited identifier cannot redirect the join.
    frozen_pubs = {row['publication_id']: row for row in original['sheets']['Publications']}
    for sid, role in original['roles'].items():
        native = role['SCREEN_ID']; link = li.get(native)
        if link is None: raise ValueError('Selected association missing: ' + native)
        old = next(row for row in original['sheets']['Screens'] if row['screen_id'] == sid)
        if old.get('publication_id'):
            pid = old['publication_id']; pub = pi[link['PUBLICATION_ID']]; identity = frozen_pubs[pid]
            field = 'pmid' if pub['SOURCE_TYPE'] == 'pubmed' else 'doi'
            if str(identity.get(field)) != str(pub['SOURCE_ID']): raise ValueError('Frozen publication identity mismatch')
            if link['PUBLICATION_ID'] in used and used[link['PUBLICATION_ID']] != pid:
                raise ValueError('Ambiguous native-to-internal publication identity')
            used[link['PUBLICATION_ID']] = pid
    if len(set(used.values())) != len(used): raise ValueError('Internal publication ID reused for distinct native records')
    for native_pid in sorted({li[r['SCREEN_ID']]['PUBLICATION_ID'] for r in original['roles'].values()}):
        pub = pi[native_pid]; pid = used.get(native_pid)
        if pid is None:
            pid = stable_id('PUB', 'ORCS publication', native_pid, pub['SOURCE_TYPE'], pub['SOURCE_ID'])
            if pid in pub_rows:
                field = 'pmid' if pub['SOURCE_TYPE']=='pubmed' else 'doi'
                if pub_rows[pid].get(field) != pub.get(field.upper()): raise ValueError('New publication ID collision')
            else:
                pub_rows[pid] = {'publication_id': pid}; p['sheets']['Publications'].append(pub_rows[pid])
            used[native_pid] = pid
        if pid not in pub_rows: raise ValueError('Original publication ID missing from workbook')
        row = pub_rows[pid]
        locator = pub['SOURCE_METADATA_FILE'] + '; SHA256=' + pub['SOURCE_METADATA_SHA256']
        eid = add_evidence('header', native_pid, 'publication', pid, locator, pub, pub['RETRIEVED_AT'],
                           'conflicting identifiers; review required' if pub['REVIEW_ISSUES'] else 'reported publication header')
        for source, dest in [('TITLE','title_original'), ('AUTHORS','author_list_ reported'), ('JOURNAL','journal_reported'), ('PMID','pmid'), ('DOI','doi')]:
            propose('Publications', row, dest, pub.get(source), locator + '; ' + pub['FIELD_LOCATORS'].get(source, source))
        try: year = date.fromisoformat(pub['PUBLICATION_DATE']).year
        except (ValueError, TypeError): year = None
        propose('Publications', row, 'publication_year', year, locator + '; PUBLICATION_DATE', 'valid ISO date -> year')
        for source, dest, prefix in [('PMID','pmid_link','https://pubmed.ncbi.nlm.nih.gov/'), ('DOI','doi_link','https://doi.org/')]:
            value = pub.get(source)
            url = prefix + quote(value, safe='/') + ('/' if source == 'PMID' else '') if value else None
            propose('Publications', row, dest, url, locator + '; ' + source, 'identifier -> derived URL; not fetched')
        propose('Publications', row, 'retrival_sources', 'BioGRID ORCS saved publication header; evidence_id=' + eid, locator, 'provenance append', True)
        propose('Publications', row, 'publication_notes', 'ORCS publication header: ' + json.dumps(pub, ensure_ascii=False), locator, 'reported metadata and provenance append', True)
        propose('Publications', row, 'evidence_status', 'publication header reported; ' + ('identifier conflict requires owner review' if pub['REVIEW_ISSUES'] else 'explicit native association verified'), locator, 'status append', True)
        if pub['REVIEW_ISSUES']: p['publication_review'][pid] = pub['REVIEW_ISSUES']
        mapping.append({'PUBLICATION_ID': native_pid, 'publication_id': pid, 'SOURCE_TYPE': pub['SOURCE_TYPE'], 'SOURCE_ID': pub['SOURCE_ID'], 'evidence_id': eid})
    association_map = []
    for sid, role in original['roles'].items():
        native = role['SCREEN_ID']; n = si[native]; link = li[native]; pid = used[link['PUBLICATION_ID']]; row = screen_rows[sid]
        locator = 'data/orcs/publication-screen-links.json; SCREEN_ID=' + native
        eid = add_evidence('association', native, 'screen', sid, locator,
                           {**link, 'publication_id': pid, 'screen_id': sid, 'relationship': 'direct publication association; no dataset attribution'},
                           association_retrieved_at or pi[link['PUBLICATION_ID']]['RETRIEVED_AT'], 'direct publication association')
        propose('Screens', row, 'publication_id', pid, locator, 'native association -> internal publication ID')
        propose('Screens', row, 'source_evidence_ids', eid, locator, 'evidence reference append', True)
        for source, dest in DIRECT.items():
            value = n.get(source); value = None if value in (None, '', '-') else value
            propose('Screens', row, dest, value, 'data/orcs/screen-index.json; SCREEN_ID=' + native + '; ' + source)
        for source, dest in [('SCREEN_FORMAT','screen_format_normalized'), ('METHODOLOGY','perturbation_type_normalized')]:
            propose('Screens', row, dest, CROSSWALKS[source].get(n.get(source)), 'data/orcs/screen-index.json; SCREEN_ID=' + native + '; ' + source, 'accepted explicit crosswalk')
        notes = {key: n[key] for key in ['CELL_TYPE','EXPERIMENTAL_SETUP','SCREEN_RATIONALE','NOTES','MOI','SCREEN_TYPE','SIGNIFICANCE_INDICATOR','SIGNIFICANCE_CRITERIA','FULL_SIZE','SCORES_SIZE','FULL_SIZE_AVAILABLE'] if n.get(key) not in (None,'','-')}
        propose('Screens', row, 'screen_notes', 'Additional reported native metadata: ' + json.dumps(notes, ensure_ascii=False), 'data/orcs/screen-index.json; SCREEN_ID=' + native, 'reported notes append', True)
        association_map.append({'SCREEN_ID': native, 'screen_id': sid, 'PUBLICATION_ID': link['PUBLICATION_ID'], 'publication_id': pid, 'evidence_id': eid, 'role_observations': role})
    for field in REVIEW:
        if field not in p['headers']['Publications']: p['headers']['Publications'].append(field)
    for row in p['sheets']['Discovery Resources']:
        propose('Discovery Resources', row, 'discovery_resources_notes',
                'Offline publication projection: 1176 selected screens; 93 supported publications; direct filter counts remain 1020 screens / 76 publications. Catalog inventories 2217 screens / 418 publications. Zero new network requests.',
                'source run summary.json; publication-index.manifest.json', 'provenance append', True)
        if row['resource_name'] == 'BioGRID ORCS':
            propose('Discovery Resources', row, 'limitations', 'Bibliographic metadata reused locally; no immune eligibility, screen grouping, dataset attribution or asset-content retrieval.', 'handoff 14', 'stage limitation append', True)
    p['native_mapping'] = {'publications': mapping, 'associations': association_map}
    p['proposed_counts'] = {name: len(rows) for name, rows in p['sheets'].items()}
    return p


def prepare(root, source, output, screen_path=None, publication_path=None, link_path=None, starting_workbook=None, legacy_schema=False):
    """Create an exclusive new run with immutable input references and a backup.

    Paths are local; fingerprint changes, corrupt receipts or altered source
    memberships fail before export. No collector/filter or network client runs.
    """
    root, source, output = Path(root).resolve(), Path(source).resolve(), Path(output).resolve()
    if output.exists(): raise FileExistsError(output)
    for name, expected in FINGERPRINTS.items():
        if sha(root / name) != expected: raise ValueError('Handoff fingerprint changed: ' + name)
    manifest = read(source / 'run_manifest.json')
    if manifest['status'] != 'complete': raise ValueError('Source run is not complete')
    paths = {'screens': Path(screen_path or root/'data/orcs/screen-index.json'),
             'publications': Path(publication_path or root/'data/orcs/publication-index.json'),
             'links': Path(link_path or root/'data/orcs/publication-screen-links.json')}
    catalogs = {key: read(path) for key, path in paths.items()}
    si, pi, li = validate_catalogs(catalogs['screens'], catalogs['publications'], catalogs['links'])
    if len(si) != 2217 or len(pi) != 418: raise ValueError('Unexpected catalog inventories')
    sm = read(root/'data/orcs/screen-index.manifest.json'); pm = read(root/'data/orcs/publication-index.manifest.json')
    if sha(paths['screens']) != sm['sha256'] or sha(paths['screens']) != pm['screen_index_sha256']:
        raise ValueError('Screen catalog manifest hash mismatch')
    for key, name in [('publications','publication-index.json'), ('links','publication-screen-links.json')]:
        if sha(paths[key]) != pm['files'][name]['sha256']: raise ValueError('Publication catalog manifest hash mismatch')
    original = read(source/'projection.json'); summary = read(source/'summary.json')
    if len(original['roles']) != 1176 or summary['matched'] != 1020 or summary['unresolved_category'] != 59 or summary['publication_siblings'] != 108 or summary['distinct_matched_publications'] != 76:
        raise ValueError('Source run counts differ from handoff')
    saved = {name: read(source/(name+'.json')) for name in ['matched_screens','unresolved_screens','publication_siblings']}
    selected = {row['SCREEN_ID'] for rows in saved.values() for row in rows}
    if selected != {r['SCREEN_ID'] for r in original['roles'].values()}: raise ValueError('Saved result selection differs from projection')
    if sum(bool(r.get('also_unresolved_category')) for r in original['roles'].values()) != 11: raise ValueError('Overlapping roles differ')
    for rows in saved.values():
        for row in rows:
            if row['native'] != si[row['SCREEN_ID']]: raise ValueError('Native screen changed since source filter')
    inputs = list(source.rglob('*')) + list((root/'data/orcs').rglob('*')) + list(paths.values()) + [root/name for name in FINGERPRINTS] + [root/'docs/orcs/enrichment-plan.md']
    refs = {str(path.resolve()): sha(path) for path in inputs if path.is_file()}
    for pub in pi.values():
        header = root/pub['SOURCE_METADATA_FILE']; receipt_path = header.with_name(pub['PUBLICATION_ID']+'.receipt.json'); receipt = read(receipt_path)
        if sha(header) != pub['SOURCE_METADATA_SHA256'] or receipt['sha256'] != pub['SOURCE_METADATA_SHA256'] or receipt['retrieved_at'] != pub['RETRIEVED_AT'] or receipt['status'] != 'success':
            raise ValueError('Saved publication evidence/receipt corrupt: ' + pub['PUBLICATION_ID'])
    wb_path = Path(starting_workbook or source/'DeathMap-AI-v1-reference-output.xlsx')
    if not starting_workbook and sha(wb_path) != manifest['workbook_sha256']: raise ValueError('Starting workbook differs from completed source run')
    refs[str(wb_path.resolve())] = sha(wb_path)
    workbook = existing_rows(wb_path); template = existing_rows(root/'data/templates/DeathMap-AI-v1-reference-output.xlsx')
    if list(workbook) != list(template): raise ValueError('Template sheet order mismatch')
    for name, value in template.items():
        if workbook[name]['headers'][:len(value['headers'])] != value['headers']: raise ValueError('Template header mismatch: ' + name)
    now = datetime.now(timezone.utc).isoformat()
    browse_receipt = read(root/'data/orcs/publication-source/20261004-v01/browse-receipt.json')
    if browse_receipt['sha256'] != pm['browse_sha256'] or browse_receipt['http_status'] != 200:
        raise ValueError('Association browse receipt mismatch')
    p = enrich(original, workbook, catalogs['screens'], catalogs['publications'], catalogs['links'], now, browse_receipt['retrieved_at'])
    if not legacy_schema:
        from .orcs_workbook_revision import revise
        p = revise(p,catalogs['screens'],workbook)
    if p['proposed_counts']['Publications'] != 93: raise ValueError('Unexpected publication reconciliation')
    mappings = field_map((root/'docs/orcs/workbook-field-source-map.md').read_text(encoding='utf-8'))
    if len(mappings) != 128: raise ValueError('Field map does not contain 128 mappings')
    if not legacy_schema:
        for mapping in mappings:
            if mapping['sheet']=='Screens' and mapping['field']=='library_gene_count_reported':
                mapping.update(classification='reported',source_rule='Owner revision WR2: native FULL_SIZE copied as integer; no independent library validation')
        mappings.append({'sheet':'Publications','field':'evidence_id_link','classification':'provenance','source_rule':'Evidence IDs separated from retrival_sources; owner revision WR1'})
        revision_doc=root/'docs/specifications/15 ORCS Workbook ID and Field Revision.md'
        refs[str(revision_doc.resolve())]=sha(revision_doc)
    audited = {(x['sheet'], x['id'], x['field']) for x in p['cell_audit']}
    for m in mappings:
        original_ids = {row[workbook[m['sheet']]['headers'][0]] for row in workbook[m['sheet']]['rows']}
        if p.get('id_migration'):
            id_map={x['old_id']:x['new_id'] for x in p['id_migration']}
            original_ids={id_map.get(identity,identity) for identity in original_ids}
        for row in p['sheets'][m['sheet']]:
            identity = row[p['headers'][m['sheet']][0]]
            if (m['sheet'], identity, m['field']) not in audited:
                populated = row.get(m['field']) not in (None,'')
                p['cell_audit'].append({'sheet':m['sheet'],'id':identity,'field':m['field'],
                    'status':('preserved existing' if identity in original_ids else 'generated evidence') if populated else 'unresolved',
                    'classification':m['classification'],'reason':m['source_rule'],
                    'source_locator':row.get('source_locator') if m['sheet']=='Sources & Evidence' else m['source_rule']})
    output.mkdir(parents=True)
    shutil.copyfile(wb_path, output/'workbook-original.xlsx')
    write(output/'projection.json', p); write(output/'native-id-mapping.json', p['native_mapping'])
    if p.get('id_migration'): write(output/'id-migration.json',p['id_migration'])
    write(output/'field-population-audit.json', {'column_mappings':mappings,'cells':p['cell_audit']})
    write(output/'input-references.json', {'sha256_by_absolute_path':refs})
    code_files = ['src/deathmap_ai/orcs_publication_enrichment.py','src/deathmap_ai/orcs_workbook_revision.py','scripts/export_orcs_filter_workbook.mjs','scripts/enrich_orcs_publications.py','src/deathmap_ai/orcs_filter_projection.py','src/deathmap_ai/orcs_filter_verify.py']
    write(output/'run_manifest.json', {'version':p['version'],'starting_workbook':str(wb_path.resolve()),'status':'pending export and verification','started_at':now,
          'source_run':str(source),'source_summary':summary,'input_paths':{k:str(v.resolve()) for k,v in paths.items()},
          'input_hashes':refs,'mapping_document_sha256':sha(root/'docs/orcs/workbook-field-source-map.md'),
          'code_hashes':{f:sha(root/f) for f in code_files},'network_policy':'offline; no filtering, collector, refresh or network calls',
          'network_requests':0,'source_screen_retrieved_at':manifest['cache_retrieved_at'],'source_reference_retrieved_at':manifest['reference_retrieval_date'],
          'publication_retrieval_range':[pm['earliest_page_retrieval'],pm['latest_page_retrieval']], 'artifact_bundle_version':'26.930.11008'})
    return p


def verify_enriched(run, root):
    """Independently reconcile exported cells, entity keys, roles and preservation.

    This read-only acceptance gate raises on any unexplained difference. It writes
    a verification report, but completion requires a separate visual-review gate.
    """
    run, root = Path(run), Path(root); p = read(run/'projection.json'); m = read(run/'run_manifest.json')
    before = load_workbook(run/'workbook-original.xlsx', data_only=False)
    after = load_workbook(run/'DeathMap-AI-v1-reference-output.xlsx', data_only=False)
    template = load_workbook(root/'data/templates/DeathMap-AI-v1-reference-output.xlsx')
    if before.sheetnames != after.sheetnames or template.sheetnames != after.sheetnames: raise ValueError('Sheet order changed')
    appends = {(x['sheet'],x['id'],x['field']):x for x in p['approved_appends']}
    changes = {(x['sheet'],x['id'],x['field']):x for x in p.get('approved_changes',[])}
    migration = {x['old_id']:x['new_id'] for x in p.get('id_migration',[])}
    counts = {}; formulas = 0
    for sheet in before:
        target = after[sheet.title]; headers = [c.value for c in target[1]]
        if headers != p['headers'][sheet.title]: raise ValueError('Export header mismatch')
        th = [c.value for c in template[sheet.title][1]]
        compared = [h for h in headers if h!='evidence_id_link']
        if compared[:len(th)] != th: raise ValueError('Template headers changed')
        for row in list(sheet.rows)[1:]:
            for cell in row:
                if cell.value is None: continue
                identity = migration.get(row[0].value,row[0].value); field = sheet.cell(1, cell.column).value
                expected = appends.get((sheet.title,identity,field),{}).get('value',cell.value)
                expected = changes.get((sheet.title,identity,field),{}).get('value',expected)
                observed = target.cell(cell.row,headers.index(field)+1)
                if observed.value != expected: raise ValueError(f'Existing cell changed: {sheet.title}!{cell.coordinate}')
                if cell.data_type == 'f': formulas += 1
                if copy.copy(cell.font) != copy.copy(observed.font) or copy.copy(cell.border) != copy.copy(observed.border) or cell.number_format != observed.number_format:
                    raise ValueError(f'Existing style changed: {sheet.title}!{cell.coordinate}')
        actual = [dict(zip(headers,row)) for row in list(target.values)[1:] if any(v is not None for v in row)]
        key = headers[0]; indexed = {row[key]:row for row in actual}
        if len(indexed) != len(actual) or len(actual) != len(p['sheets'][sheet.title]): raise ValueError('Row count/identity mismatch')
        for row in p['sheets'][sheet.title]:
            for field, value in row.items():
                observed = indexed[row[key]].get(field)
                if field == 'retrieved_at' and isinstance(observed,datetime):
                    expected = datetime.fromisoformat(value).replace(tzinfo=None)
                    if abs((observed-expected).total_seconds()) < .002: continue
                if observed != value and not (observed in (None,'') and value in (None,'')):
                    raise ValueError(f'Projection cell mismatch: {sheet.title}/{row[key]}/{field}')
        for row in target:
            for cell in row:
                if cell.data_type == 'e': raise ValueError('Spreadsheet formula error: ' + cell.coordinate)
        counts[sheet.title] = len(actual)
    source = read(Path(m['source_run'])/'projection.json')
    expected_roles = {migration.get(k,k):v for k,v in source['roles'].items()}
    if p['roles'] != expected_roles: raise ValueError('Filter roles changed')
    if counts['Screens'] != 1176 or counts['Publications'] != 93: raise ValueError('Unexpected entity counts')
    for name in ['Screen Groups','Datasets','Screen-Dataset Links']:
        if p['sheets'][name] != existing_rows(run/'workbook-original.xlsx')[name]['rows']: raise ValueError('Unsupported entities added')
    pub_ids = {r['publication_id'] for r in p['sheets']['Publications']}; screen_ids = {r['screen_id'] for r in p['sheets']['Screens']}
    if not {migration.get(r['publication_id'],r['publication_id']) for r in source['sheets']['Publications']} <= pub_ids: raise ValueError('Historical publication mapping changed')
    evidence = {r['evidence_id']:r for r in p['sheets']['Sources & Evidence']}
    for row in evidence.values():
        typ = row['supports_entity_type']
        if typ in ['publication','screen'] and row['supports_entity_id'] not in (pub_ids if typ == 'publication' else screen_ids): raise ValueError('Evidence foreign key invalid')
    for row in p['sheets']['Screens']:
        if row.get('publication_id') not in pub_ids: raise ValueError('Screen publication foreign key invalid')
        for eid in re.split(r';\s*|\n',row['source_evidence_ids']):
            if eid not in evidence: raise ValueError('Screen evidence reference missing')
    pubmap = {r['PUBLICATION_ID']:r['publication_id'] for r in p['native_mapping']['publications']}
    screenmap = {r['SCREEN_ID']:r for r in p['native_mapping']['associations']}
    catalogs = {key:read(Path(filename)) for key,filename in m['input_paths'].items()}
    si, pi, li = validate_catalogs(catalogs['screens'],catalogs['publications'],catalogs['links'])
    expected_native_pubs = {li[r['SCREEN_ID']]['PUBLICATION_ID'] for r in source['roles'].values()}
    if set(pubmap) != expected_native_pubs or len(set(pubmap.values())) != len(pubmap): raise ValueError('Publication mapping reconciliation failed')
    for native, assoc in screenmap.items():
        if assoc['PUBLICATION_ID'] != li[native]['PUBLICATION_ID'] or assoc['publication_id'] != pubmap[assoc['PUBLICATION_ID']]: raise ValueError('Association mapping differs from source')
    # Recompute proposals from immutable inputs rather than accepting a workbook
    # and its projection agreeing with each other as sufficient source evidence.
    browse = read(root/'data/orcs/publication-source/20261004-v01/browse-receipt.json')
    recomputed = enrich(source,existing_rows(run/'workbook-original.xlsx'),catalogs['screens'],catalogs['publications'],catalogs['links'],m['started_at'],browse['retrieved_at'])
    if migration:
        from .orcs_workbook_revision import revise
        recomputed = revise(recomputed,catalogs['screens'],existing_rows(run/'workbook-original.xlsx'))
        for sheet,key,prefix,width in [('Publications','publication_id','PUB',4),('Screens','screen_id','SCR',4),('Sources & Evidence','evidence_id','EVI',5)]:
            ids=[r[key] for r in p['sheets'][sheet]]
            if len(set(ids))!=len(ids) or any(not re.fullmatch(prefix+r'-\d{'+str(width)+'}',x) for x in ids): raise ValueError('Invalid sequential IDs')
        h=p['headers']['Publications']
        if h.index('evidence_id_link')!=h.index('retrival_sources')+1: raise ValueError('Evidence column misplaced')
        for row in p['sheets']['Publications']:
            if 'evidence_id=' in row['retrival_sources']: raise ValueError('Evidence ID remains in retrieval sources')
            for eid in (row.get('evidence_id_link') or '').split('; '):
                if eid and (eid not in evidence or evidence[eid]['supports_entity_id']!=row['publication_id'] or evidence[eid]['supports_entity_type']!='publication'): raise ValueError('Publication evidence link invalid')
    if recomputed['sheets'] != p['sheets'] or recomputed['native_mapping'] != p['native_mapping'] or recomputed['approved_appends'] != p['approved_appends']:
        raise ValueError('Projection does not reproduce from accepted input evidence')
    if p['cell_audit'][:len(recomputed['cell_audit'])] != recomputed['cell_audit']:
        raise ValueError('Population audit does not reproduce')
    if screenmap['2469']['publication_id'] != pubmap['1165']: raise ValueError('Prepub association missing')
    prepub = next(r for r in p['sheets']['Publications'] if r['publication_id'] == pubmap['1165'])
    if prepub.get('pmid') not in (None,''): raise ValueError('Prepub page PMID promoted')
    for sheet_name in ['Screens','Publications']:
        sheet = after[sheet_name]; headers = [c.value for c in sheet[1]]
        if headers[-2:] != REVIEW: raise ValueError('Review columns not rightmost')
        for row in list(sheet.rows)[1:]:
            attention = sheet_name == 'Screens' and row[0].value in p['roles'] or sheet_name == 'Publications' and row[0].value in p['publication_review']
            color = 'F4B183' if attention and row[-2].value in (None,'') else 'FFF2CC'
            if any(not str(c.fill.fgColor.rgb).endswith(color) for c in row[-2:]): raise ValueError('Reviewer color mismatch')
    for filename, digest in m['input_hashes'].items():
        if sha(Path(filename)) != digest: raise ValueError('Input/history changed: ' + filename)
    if sha(run/'workbook-original.xlsx') != sha(Path(m.get('starting_workbook',str(Path(m['source_run'])/'DeathMap-AI-v1-reference-output.xlsx')))): raise ValueError('Backup mismatch')
    for filename, digest in m['code_hashes'].items():
        if sha(root/filename) != digest: raise ValueError('Implementation changed during run: ' + filename)
    artifact_checks = read(run/'artifact-checks.json')
    if 'Cell search matched 0 entries.' not in artifact_checks['errors']:
        raise ValueError('Artifact formula error scan requires investigation')
    result = {'status':'passed','verified_at':datetime.now(timezone.utc).isoformat(),'row_counts':counts,
              'existing_formulas_preserved':formulas,'input_files_unchanged':len(m['input_hashes']),
              'network_requests':0,'roles_and_ids_preserved':True,'template_headers_and_review_fields_verified':True,
              'workbook_sha256':sha(run/'DeathMap-AI-v1-reference-output.xlsx')}
    write(run/'verification.json',result)
    return result


def complete(run, root):
    """Mark complete only with passing verification and a hashed visual review.

    The visual receipt is written after inspecting representative publication,
    screen, evidence and reviewer images. It is explicit review evidence rather
    than treating successful rendering alone as a visual acceptance decision.
    """
    from collections import Counter
    run = Path(run); visual = read(run/'visual-verification.json')
    if visual.get('status') != 'passed': raise ValueError('Visual review not passed')
    if set(visual.get('reviewed_ranges',[])) != {'publications','screens','evidence','reviewers'}:
        raise ValueError('Required visual review ranges missing')
    for filename, digest in visual['preview_hashes'].items():
        if sha(run/filename) != digest: raise ValueError('Reviewed preview changed')
    report = verify_enriched(run,root); p = read(run/'projection.json'); manifest = read(run/'run_manifest.json')
    statuses = dict(Counter(a['status'] for a in p['cell_audit']))
    fields = dict(Counter(a['sheet']+'/'+a['field'] for a in p['cell_audit'] if a['status']=='added'))
    summary = {'status':'complete','row_counts':report['row_counts'],'direct_screen_hits':1020,'direct_hit_publications':76,
               'unresolved_category_records':59,'publication_context_records':108,'unresolved_context_overlaps':11,
               'field_population_statuses':statuses,'fields_added':fields,'publication_identifier_conflicts':p['publication_review'],
               'original_ids_preserved':True,'network_requests':0,'offline_tests_passed':visual.get('offline_tests_passed'),
               'limitations':['Two linked publications have no reported journal; cells left blank.',
                   'Screen 2469 remains unresolved-category; publication 1165 page number 28 is not promoted to PMID.',
                   'No immune eligibility, grouping or exact dataset attribution established.' + (' FULL_SIZE output mapping is owner-approved; library size is not independently validated.' if p.get('id_migration') else ' Short sequential IDs remain pending RD1.')],
               'learning_checkpoint':'Explicit native associations add bibliography and missing publication links while preserving scientific filter outcomes.'}
    write(run/'summary.json',summary)
    if p.get('id_migration'): summary['id_migration_records']=len(p['id_migration']); summary['original_ids_preserved']=False; summary['original_identity_mapping_preserved']=True; write(run/'summary.json',summary)
    (run/'summary.md').write_text(
        '# ORCS publication-enriched workbook\n\nComplete offline projection: 1,176 screens, 93 publications and 3,618 evidence records. '
        'The original 1,020 direct screen hits and 76 direct-hit publications are unchanged.\n\n'
        'Added 93 titles, 93 author strings, 93 years, 91 journals, one DOI/URL and the missing screen 2469 publication link. '
        'Two journals remain blank. 59 unresolved-category records remain unresolved; 11 overlap with publication context.\n\n'
        'Publication 1165 retains its prepub DOI and conflicting page identifier; PMID 28 is not populated. '
        'Reviewer fields are rightmost, yellow, and orange where owner attention is requested. '
        'No new groups, datasets or screen-dataset links. Zero new network requests.\n\n'
        'Learning checkpoint: local bibliography enrichment adds metadata and direct associations without changing biological eligibility.\n',encoding='utf-8')
    if p.get('id_migration'):
        (run/'summary.md').write_text('# ORCS workbook ID and field revision\n\n'
            'Complete offline rerun: 93 publications, 1,176 screens and 3,618 evidence records. '
            'Sequential IDs use PUB-0001, SCR-0001 and EVI-00001 formats; id-migration.json preserves every prior identity.\n\n'
            'Publication evidence IDs are separated into evidence_id_link immediately after retrival_sources. '
            'FULL_SIZE populates library_gene_count_reported under the owner-approved mapping. Native values and provenance remain; '
            'this mapping does not independently validate experimental library size.\n\n'
            'The 1,020 direct hits, 76 direct-hit publications, 59 unresolved-category records and 11 unresolved/context overlaps remain unchanged. '
            'Two journals remain blank; publication 1165 retains its conflict and no PMID 28 is populated. '
            'No new dataset/group rows or network requests. Reviewer entries remain intact.\n\n'
            'Learning checkpoint: a reversible ID registry keeps short workbook identifiers aligned with all foreign keys and original native identities.\n',encoding='utf-8')
    manifest.update(status='complete',completed_at=datetime.now(timezone.utc).isoformat(),workbook_sha256=report['workbook_sha256'],
                    offline_tests_passed=visual.get('offline_tests_passed'),visual_verification=visual)
    write(run/'run_manifest.json',manifest)
    return summary


def main():
    """Prepare or verify an explicitly selected offline publication-enrichment run."""
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('source',type=Path); parser.add_argument('output',type=Path)
    parser.add_argument('--root',type=Path,default=Path(__file__).resolve().parents[2])
    parser.add_argument('--screens',type=Path); parser.add_argument('--publications',type=Path); parser.add_argument('--links',type=Path)
    parser.add_argument('--verify',action='store_true')
    parser.add_argument('--complete',action='store_true')
    parser.add_argument('--starting-workbook',type=Path)
    parser.add_argument('--legacy-schema',action='store_true',help='Retain handoff 14 schema for historical reproduction')
    args = parser.parse_args()
    if args.complete: print(json.dumps(complete(args.output,args.root),indent=2))
    elif args.verify: print(json.dumps(verify_enriched(args.output,args.root),indent=2))
    else:
        p = prepare(args.root,args.source,args.output,args.screens,args.publications,args.links,args.starting_workbook,args.legacy_schema)
        print(json.dumps({'output':str(args.output),'status':'pending export and verification','counts':p['proposed_counts']},indent=2))

if __name__ == '__main__': main()
