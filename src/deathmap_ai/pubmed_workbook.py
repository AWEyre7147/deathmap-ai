"""Export PMID metadata into the finalized owner template with literal dates.

Only sheet data, dimensions and new body styles change. Keeping original XML
envelopes preserves Excel's compatibility-prefix declarations and template
formatting, including column widths, frozen panes and header styles.
"""
from collections import defaultdict
from datetime import date
from pathlib import Path
import copy
import io
import re
import zipfile
import xml.etree.ElementTree as ET
import openpyxl
from openpyxl.utils import get_column_letter

NS = 'http://schemas.openxmlformats.org/spreadsheetml/2006/main'
MC = 'http://schemas.openxmlformats.org/markup-compatibility/2006'
ET.register_namespace('', NS)
q = lambda name: '{' + NS + '}' + name


def date_text(value):
    """Convert ISO source dates to literal owner-selected review text.

    Parameters
    ----------
    value : str
        ISO day/month/year precision or an unchanged reported date string.

    Returns
    -------
    str
        MM;DD;YYYY, MM;YYYY, or the original year/free-text value. Missing
        components are never inferred; this is not an Excel numeric date.
    """
    if re.fullmatch(r'\d{4}-\d{2}-\d{2}', value):
        return date.fromisoformat(value).strftime('%m;%d;%Y')
    if re.fullmatch(r'\d{4}-\d{2}', value):
        year, month = value.split('-')
        return f'{month};{year}'
    return value


def subset_rows(database, pmids):
    """Select publications and recompute MeSH counts for a PMID subset.

    Parameters
    ----------
    database : dict
        Canonical enrichment rows with ISO dates and descriptor/qualifier pairs.
    pmids : list of str
        Exact publication identifiers for this output.

    Returns
    -------
    dict
        Copied workbook rows with text dates and subset-specific MeSH lists.
        The input database is not mutated.
    """
    selected = set(pmids)
    rows = {name: [list(row) for row in values if row[0] in selected]
            for name, values in database['rows'].items() if name != 'MeSH'}
    labels = {(r[0], r[2]): (r[1], r[3]) for r in database['rows']['MeSH']}
    groups = defaultdict(set)
    for row in rows['MeSH Links']:
        groups[(row[1], row[2])].add(row[0])
    rows['MeSH'] = [[key[0], labels[key][0], key[1], labels[key][1], len(values), '; '.join(sorted(values, key=int))]
                    for key, values in sorted(groups.items())]
    for record in rows['Publication']:
        record[4] = date_text(record[4])
    return rows


def replace_block(xml, tag, element):
    """Replace one XML block while preserving its surrounding namespace envelope.

    Parameters
    ----------
    xml : str
        Original workbook part.
    tag : str
        Single element name to replace.
    element : Element
        Updated element subtree.

    Returns
    -------
    str
        Original outer XML with one replaced subtree.

    Raises
    ------
    AssertionError
        The expected element block is absent.
    """
    prefix = r'(?:[A-Za-z_][\w.-]*:)?'
    result, count = re.subn(r'<' + prefix + tag + r'\b[^>]*>.*?</' + prefix + tag + '>',
                            lambda _: ET.tostring(element, encoding='unicode'), xml, count=1, flags=re.S)
    assert count == 1, tag
    return result


def validate_namespaces(payload):
    """Reject compatibility declarations with undeclared XML prefixes.

    Parameters
    ----------
    payload : bytes
        Workbook XML part with native compatibility declarations.

    Returns
    -------
    None
        No mutation; raises when the scope check fails.

    Raises
    ------
    AssertionError
        An mc:Ignorable prefix is missing from its namespace scope.
    """
    scopes, pending = [], []
    for event, item in ET.iterparse(io.BytesIO(payload), events=('start-ns', 'start', 'end')):
        if event == 'start-ns':
            pending.append(item)
        elif event == 'start':
            scope = dict(scopes[-1]) if scopes else {}
            scope.update(pending)
            pending = []
            assert all(prefix in scope for prefix in item.get('{' + MC + '}Ignorable', '').split())
            scopes.append(scope)
        else:
            scopes.pop()


def export_workbook(template, output, rows):
    """Write and verify a new workbook from the finalized owner template.

    Parameters
    ----------
    template : Path or str
        Header-only finalized PubMed template.
    output : Path or str
        New workbook destination; parent directories are created as needed.
    rows : dict
        Reviewed columns per sheet, including literal text publication dates.

    Returns
    -------
    None
        Writes the workbook and compares all projected rows after reopening it.

    Raises
    ------
    FileExistsError
        Destination already exists; owner workbooks cannot be overwritten.
    AssertionError
        Template, namespace, package preservation or row comparison fails.
    """
    output = Path(output)
    if output.exists():
        raise FileExistsError(output)
    with zipfile.ZipFile(template) as original:
        parts = {name: original.read(name) for name in original.namelist()}
    styles = ET.fromstring(parts['xl/styles.xml'])
    xfs = styles.find(q('cellXfs'))
    body_id = len(xfs)
    # Body formatting follows the template's default style, rather than a font
    # index or wrapped/tall-row presentation inherited from a retired template.
    body = copy.deepcopy(xfs[0])
    xfs.append(body)
    text_id = len(xfs)
    text_style = copy.deepcopy(body)
    text_style.set('numFmtId', '49')
    text_style.set('applyNumberFormat', '1')
    xfs.append(text_style)
    xfs.set('count', str(len(xfs)))
    parts['xl/styles.xml'] = replace_block(parts['xl/styles.xml'].decode(), 'cellXfs', xfs).encode()
    workbook = ET.fromstring(parts['xl/workbook.xml'])
    rels = {r.get('Id'): r.get('Target') for r in ET.fromstring(parts['xl/_rels/workbook.xml.rels'])}
    changed = {'xl/styles.xml'}
    for sheet in workbook.find(q('sheets')):
        name = sheet.get('name')
        target = rels[sheet.get('{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id')]
        member = target.lstrip('/') if target.startswith('/') else 'xl/' + target
        xml = parts[member].decode()
        document = ET.fromstring(xml)
        data = document.find(q('sheetData'))
        assert len(data) == 1, 'Expected header-only template'
        for number, values in enumerate(rows[name], 2):
            row = ET.SubElement(data, q('row'), {'r': str(number)})
            for index, value in enumerate(values, 1):
                cell = ET.SubElement(row, q('c'), {'r': get_column_letter(index) + str(number), 's': str(body_id)})
                if isinstance(value, int):
                    ET.SubElement(cell, q('v')).text = str(value)
                else:
                    cell.set('t', 'inlineStr')
                    if name == 'Publication' and index == 5:
                        cell.set('s', str(text_id))
                    ET.SubElement(ET.SubElement(cell, q('is')), q('t')).text = str(value)
        xml = replace_block(xml, 'sheetData', data)
        extent = 'A1:' + get_column_letter(len(data[0])) + str(len(rows[name]) + 1)
        dimension = '<dimension xmlns="' + NS + '" ref="' + extent + '"/>'
        xml, count = re.subn(r'<(?:[A-Za-z_][\w.-]*:)?dimension\b[^>]*/>', lambda _: dimension, xml, count=1)
        # Some valid header-only exports omit dimension. Insert it before views,
        # after any sheet properties, without reserializing the owner XML envelope.
        if count == 0:
            xml, count = re.subn(r'(<(?:[A-Za-z_][\w.-]*:)?sheetViews\b)', lambda m: dimension + m[1], xml, count=1)
        assert count == 1
        parts[member] = xml.encode()
        changed.add(member)
    for name, payload in parts.items():
        if name.endswith('.xml'):
            validate_namespaces(payload)
    output.parent.mkdir(parents=True, exist_ok=True)
    with zipfile.ZipFile(output, 'w', zipfile.ZIP_DEFLATED) as archive:
        for name, payload in parts.items():
            archive.writestr(name, payload)
    # Inspect the independently reopened package, including every projected row.
    saved = openpyxl.load_workbook(output, read_only=True)
    for name, values in rows.items():
        actual = [tuple('' if v is None else v for v in row) for row in list(saved[name].values)[1:]]
        assert actual == [tuple(row) for row in values], name
    saved.close()
    with zipfile.ZipFile(template) as original, zipfile.ZipFile(output) as result:
        for name in original.namelist():
            if name not in changed:
                assert original.read(name) == result.read(name), name
            elif name.startswith('xl/worksheets/'):
                before, after = ET.fromstring(original.read(name)), ET.fromstring(result.read(name))
                for tag in ['cols', 'sheetViews', 'sheetFormatPr', 'pageMargins']:
                    assert ET.tostring(before.find(q(tag))) == ET.tostring(after.find(q(tag)))
                assert ET.tostring(before.find(q('sheetData'))[0]) == ET.tostring(after.find(q('sheetData'))[0])
